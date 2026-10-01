"""Offline exporter tests. All papers, quotes and validated rows here are fixtures.

These payloads deliberately stand in for the core's PDF/hash/quote validation;
none of the fixture metadata represents a searched paper or scientific result.
"""
from __future__ import annotations

import builtins
import copy
import http.client
import io
import os
import socket
import stat
import subprocess
import tempfile
import urllib.request
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml.ns import qn

from paperflow.engine.journals.models import JournalError
from paperflow.engine.literature import writing_export


@pytest.fixture(autouse=True)
def block_word_network_and_models(monkeypatch):
    """Any external/Word/LLM attempt fails, rather than silently making a request."""
    def forbidden(*args, **kwargs):
        raise AssertionError("offline fixture: network, Word and LLM are forbidden")

    monkeypatch.setattr(socket.socket, "connect", forbidden)
    monkeypatch.setattr(socket.socket, "connect_ex", forbidden)
    monkeypatch.setattr(socket, "create_connection", forbidden)
    monkeypatch.setattr(urllib.request, "urlopen", forbidden)
    monkeypatch.setattr(http.client.HTTPConnection, "connect", forbidden)
    monkeypatch.setattr(http.client.HTTPSConnection, "connect", forbidden)
    monkeypatch.setattr(subprocess, "Popen", forbidden)
    monkeypatch.setattr(os, "startfile", forbidden, raising=False)
    original_import = builtins.__import__
    blocked = ("win32com", "comtypes", "pythoncom", "openai", "anthropic", "litellm", "paperflow.engine.live_bridge")

    def guarded_import(name, *args, **kwargs):
        if any(name == prefix or name.startswith(prefix + ".") for prefix in blocked):
            forbidden()
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", guarded_import)


@pytest.fixture
def validated_payload():
    """Explicit validated-payload fixture; no file existence is implied."""
    def row(citation_id, paper_id, reading_id, sha, page, fragment, kind="author_claim"):
        return {
            "fixture": True,
            "citation_id": citation_id, "paper_id": paper_id, "reading_id": reading_id,
            "claim_index": 0, "file_sha256": sha,
            "section": "Fixture source Methods (not a draft chapter)",
            "kind": kind, "text": "Fixture reading claim; not a scientific assertion.",
            "evidence": [{"fragment_id": fragment, "page_number": page,
                          "quote": "FIXTURE SOURCE QUOTE THAT MUST NOT BE COPIED"}],
        }

    return {
        "fixture": True,
        "project_id": "writing-fixture", "revision": 2,
        "profile": {
            "title": "离线草稿测试 — Unicode α", "research_question": "如何定位 fixture 文献证据？",
            "language": "zh", "own_materials": [{"id": "UM1", "text": "Fixture supplied observation."}],
        },
        "draft": {
            "title": "文献支持草稿 fixture — 中文与 Ω", "language": "zh",
            "outline": [
                {"id": "related", "title": "相关工作", "level": 1, "evidence_ids": ["cite-alpha", "cite-beta"]},
                {"id": "results", "title": "结果（待补充）", "level": 1, "evidence_ids": []},
            ],
            "sections": [
                {"id": "related", "title": "相关工作", "level": 1, "paragraphs": [
                    {"text": "第一篇 fixture 的作者主张重述。", "kind": "literature_summary", "citation_ids": ["cite-alpha"], "own_material_ids": []},
                    {"text": "多文献推断 fixture，不冒充原文。", "kind": "literature_inference", "citation_ids": ["cite-beta", "cite-alpha-again", "cite-alpha"], "own_material_ids": []},
                    {"text": "建议开展对照实验；尚未开展。", "kind": "author_proposal", "citation_ids": [], "own_material_ids": []},
                    {"text": "用户实际提供的 fixture 材料。", "kind": "user_material", "citation_ids": [], "own_material_ids": ["UM1"]},
                ]},
                {"id": "results", "title": "结果（待补充）", "level": 1, "paragraphs": [
                    {"text": "", "kind": "placeholder", "citation_ids": [], "own_material_ids": []},
                ]},
            ],
        },
        "citations": {
            "cite-alpha": row("cite-alpha", "paper-fixture-alpha", "reading-fixture-alpha", "a" * 64, 3, "fragment-fixture-a3"),
            "cite-beta": row("cite-beta", "paper-fixture-beta", "reading-fixture-beta", "b" * 64, 7, "fragment-fixture-b7", "agent_inference"),
            "cite-alpha-again": row("cite-alpha-again", "paper-fixture-alpha", "reading-fixture-alpha", "a" * 64, 4, "fragment-fixture-a4"),
        },
        "references": [
            {"fixture": True, "paper_id": "paper-fixture-alpha", "title": "Fixture Alpha 中文研究",
             "authors": ["陈测试", "Zoë Ångström"], "year": 2021, "venue": "Fixture Journal",
             "doi": "10.0000/fixture-alpha", "arxiv_id": "2101.00000v2",
             "landing_url": "https://example.org/fixture-alpha"},
            {"fixture": True, "paper_id": "paper-fixture-beta", "title": "Fixture Beta — résumé",
             "authors": [{"given": "Ada", "family": "Fixture"}], "year": 2022,
             "venue": "Fixture Proceedings", "doi": "10.0000/fixture-beta",
             "landing_url": "https://example.org/fixture-beta"},
        ],
        "gaps": [{"code": "FIXTURE_NO_EXPERIMENT", "message": "尚无实测结果；只提供文献证据。"}],
        "coverage": {"backend_calls_llm": False, "evidence_check_only": True, "semantic_correctness_verified": False, "fixture": True},
    }


def rendered(payload):
    blob = writing_export.render_writing_docx(payload)
    assert isinstance(blob, bytes) and blob[:2] == b"PK"
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    return Document(io.BytesIO(blob)), xml


def all_text(document):
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(paragraphs + cells)


def single_paper_payload(payload):
    payload = copy.deepcopy(payload)
    payload["citations"] = {"cite-alpha": payload["citations"]["cite-alpha"]}
    payload["references"] = payload["references"][:1]
    payload["draft"]["outline"] = payload["draft"]["outline"][:1]
    payload["draft"]["outline"][0]["evidence_ids"] = ["cite-alpha"]
    payload["draft"]["sections"] = payload["draft"]["sections"][:1]
    payload["draft"]["sections"][0]["paragraphs"] = payload["draft"]["sections"][0]["paragraphs"][:1]
    return payload


def test_body_numbers_follow_first_paper_appearance_and_repeat_consistently(validated_payload):
    document, xml = rendered(validated_payload)
    body = [paragraph.text for paragraph in document.paragraphs]
    first = next(text for text in body if "第一篇 fixture" in text)
    inference = next(text for text in body if "多文献推断 fixture" in text)
    assert first.endswith("[1]")
    assert inference.endswith("[2][1]")
    assert not inference.endswith("[2][1][1]")
    bibliography = [text for text in body if text.startswith(("[1] ", "[2] "))]
    assert len(bibliography) == 2
    assert bibliography[0].startswith("[1] Fixture Alpha")
    assert bibliography[1].startswith("[2] Fixture Beta")
    assert "<w:instrText" not in xml and "<w:fldChar" not in xml and "<w:fldSimple" not in xml
    assert "ZOTERO_ITEM" not in xml and "ZOTERO_BIBL" not in xml


def test_first_appearance_is_not_reference_metadata_order(validated_payload):
    first, second = validated_payload["draft"]["sections"][0]["paragraphs"][:2]
    first["citation_ids"] = ["cite-alpha-again"]
    second["citation_ids"] = ["cite-beta"]
    validated_payload["draft"]["sections"][0]["paragraphs"][:2] = [second, first]
    document, _ = rendered(validated_payload)
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert next(text for text in texts if "多文献推断 fixture" in text).endswith("[1]")
    assert next(text for text in texts if "第一篇 fixture" in text).endswith("[2]")
    assert next(text for text in texts if text.startswith("[1] ")).startswith("[1] Fixture Beta")


def test_unicode_and_actual_reference_fields_are_present(validated_payload):
    document, _ = rendered(validated_payload)
    text = all_text(document)
    for expected in ("中文与 Ω", "陈测试", "Zoë Ångström", "Ada Fixture", "2021", "Fixture Journal",
                     "DOI: 10.0000/fixture-alpha", "arXiv: 2101.00000v2", "URL: https://example.org/fixture-alpha"):
        assert expected in text
    assert "可修改草稿" in text and "通用学术惯例" in text and "不代表已符合" in text
    assert "不验证用户材料" in text and "不宣称已经读懂全文" in text
    assert document.styles["Normal"].font.size.pt == 12
    assert document.styles["Normal"]._element.xml.find("宋体") >= 0
    assert document.sections[0].page_width.cm == pytest.approx(21, abs=0.02)


def test_evidence_locators_and_three_line_table_without_source_quotes(validated_payload):
    document, xml = rendered(validated_payload)
    assert len(document.tables) == 1
    table = document.tables[0]
    assert [cell.text for cell in table.rows[0].cells] == ["引用ID", "文献标题", "阅读卡ID", "PDF SHA256", "物理页码", "片段ID", "证据类别"]
    assert len(table.rows) == 4
    first = [cell.text for cell in table.rows[1].cells]
    assert first[:6] == ["cite-alpha", "Fixture Alpha 中文研究", "reading-fixture-alpha", "a" * 64, "3", "fragment-fixture-a3"]
    text = all_text(document)
    assert "reading-fixture-beta" in text and "b" * 64 in text and "fragment-fixture-a4" in text
    assert "作者主张" in text and "Agent 推断／非作者原述" in text
    assert "FIXTURE SOURCE QUOTE" not in text and "Fixture reading claim" not in text
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    assert borders.find(qn("w:top")).get(qn("w:sz")) == "12"
    assert borders.find(qn("w:bottom")).get(qn("w:sz")) == "12"
    assert borders.find(qn("w:insideV")).get(qn("w:val")) == "none"
    header_border = table.rows[0].cells[0]._tc.get_or_add_tcPr().find(qn("w:tcBorders"))
    assert header_border.find(qn("w:bottom")).get(qn("w:sz")) == "6"
    assert "物理页码" in xml


def test_all_paragraph_kinds_are_explicitly_labelled(validated_payload):
    document, _ = rendered(validated_payload)
    text = all_text(document)
    for label in ("【文献作者主张综述】", "【文献推断／非作者原述，待语义核查】",
                  "【作者建议／待验证，非已完成实验】", "【用户材料／未独立核验】", "【占位／待补充，非已验证结果】"):
        assert label in text
    assert "待补充。" in text and "尚无实测结果" in text


def test_unsaved_outline_chapter_is_a_placeholder_and_outline_only_evidence_is_not_numbered(validated_payload):
    validated_payload["draft"]["sections"] = validated_payload["draft"]["sections"][:1]
    validated_payload["draft"]["sections"][0]["paragraphs"] = validated_payload["draft"]["sections"][0]["paragraphs"][:1]
    document, _ = rendered(validated_payload)
    text = all_text(document)
    assert "本章节尚未提供正文" in text
    assert "cite-beta" in text and "仅大纲素材" in text
    assert not any(paragraph.text.startswith("[2] ") for paragraph in document.paragraphs)


def test_outline_only_draft_does_not_invent_references(validated_payload):
    validated_payload["draft"]["sections"] = []
    document, _ = rendered(validated_payload)
    text = all_text(document)
    assert "暂无正文编号引用" in text and "仅大纲素材" in text
    assert not any(paragraph.text.startswith("[1] ") for paragraph in document.paragraphs)


def test_missing_metadata_fields_are_not_filled_in(validated_payload):
    payload = single_paper_payload(validated_payload)
    payload["references"] = [{"paper_id": "paper-fixture-alpha", "title": "Title-only fixture"}]
    document, _ = rendered(payload)
    reference = next(paragraph.text for paragraph in document.paragraphs if paragraph.text.startswith("[1] "))
    assert reference == "[1] Title-only fixture"
    assert "n.d." not in reference and "Unknown" not in reference and "DOI" not in reference


@pytest.mark.parametrize("shape", ["id", "wrapped"])
def test_compatibility_reference_id_shapes(validated_payload, shape):
    for index, reference in enumerate(validated_payload["references"]):
        paper_id = reference.pop("paper_id")
        if shape == "id":
            reference["id"] = paper_id
        else:
            validated_payload["references"][index] = {"paper_id": paper_id, "paper": reference}
    rendered(validated_payload)


def test_english_output_and_unicode(validated_payload):
    validated_payload["profile"]["language"] = validated_payload["draft"]["language"] = "en"
    document, _ = rendered(validated_payload)
    text = all_text(document)
    assert "Editable literature-supported draft" in text
    assert "[Placeholder / incomplete, not verified results]" in text
    assert "Agent inference / not the authors' claim" in text
    assert "中文与 Ω" in text and "Zoë Ångström" in text
    assert document.styles["Normal"]._element.xml.find("Times New Roman") >= 0


def test_render_is_diskless_does_not_call_builder_save_or_change_payload(validated_payload, monkeypatch):
    original = copy.deepcopy(validated_payload)

    def no_disk(*args, **kwargs):
        raise AssertionError("render must not create output files or directories")

    monkeypatch.setattr(writing_export.AcademicDocxBuilder, "save", no_disk)
    monkeypatch.setattr(tempfile, "mkstemp", no_disk)
    monkeypatch.setattr(tempfile, "TemporaryDirectory", no_disk)
    monkeypatch.setattr(os, "mkdir", no_disk)
    monkeypatch.setattr(os, "makedirs", no_disk)
    monkeypatch.setattr(Path, "write_bytes", no_disk)
    monkeypatch.setattr(Path, "write_text", no_disk)
    original_open = builtins.open

    def read_only_open(file, mode="r", *args, **kwargs):
        if any(flag in mode for flag in ("w", "a", "x", "+")):
            no_disk()
        return original_open(file, mode, *args, **kwargs)

    monkeypatch.setattr(builtins, "open", read_only_open)
    rendered(validated_payload)
    assert validated_payload == original


def test_raw_source_paths_credentials_and_claimed_semantic_verification_are_not_leaked(validated_payload):
    secret = "FIXTURE-SOURCE-SECRET"
    source_path = "C:\\private-fixture\\paper.pdf"
    validated_payload["references"][0]["acquisition"] = {"file_path": source_path, "token": secret}
    validated_payload["references"][0]["raw"] = {"password": secret}
    validated_payload["citations"]["cite-alpha"]["file_path"] = source_path
    validated_payload["coverage"].update({"file_path": source_path, "token": secret, "semantic_correctness_verified": True})
    validated_payload["gaps"].append({"message": f"待定位 {source_path} token={secret}", "source_path": source_path})
    document, xml = rendered(validated_payload)
    text = all_text(document)
    assert secret not in text and source_path not in text and secret not in xml
    assert "<本地路径已省略>" in text and "<凭据已省略>" in text
    assert "不验证用户材料" in text


@pytest.mark.parametrize("manual", ["[@FAKE]", "@FakeKey", "[1]", "[1, 2]", "[2–4]", "［１２］", "ADDIN ZOTERO_ITEM fake"])
def test_reject_manual_citation_bypass(validated_payload, manual):
    validated_payload["draft"]["sections"][0]["paragraphs"][0]["text"] += manual
    with pytest.raises(JournalError):
        writing_export.render_writing_docx(validated_payload)


def test_regular_email_is_not_a_zotero_key(validated_payload):
    validated_payload["draft"]["sections"][0]["paragraphs"][2]["text"] += " fixture@example.org"
    rendered(validated_payload)


@pytest.mark.parametrize("target", ["body", "outline"])
def test_unknown_citation_id_is_rejected(validated_payload, target):
    if target == "body":
        validated_payload["draft"]["sections"][0]["paragraphs"][0]["citation_ids"] = ["cite-missing"]
    else:
        validated_payload["draft"]["outline"][0]["evidence_ids"] = ["cite-missing"]
    with pytest.raises(JournalError) as error:
        writing_export.render_writing_docx(validated_payload)
    assert error.value.code == "UNKNOWN_CITATION"


@pytest.mark.parametrize("mutation", ["missing_reference", "missing_identity", "different_identity", "nested_paper", "row_id", "reference_alias"])
def test_missing_reference_and_cross_paper_identity_rejected(validated_payload, mutation):
    if mutation == "missing_reference":
        validated_payload["references"] = validated_payload["references"][:1]
    elif mutation == "missing_identity":
        del validated_payload["references"][0]["paper_id"]
    elif mutation == "different_identity":
        validated_payload["references"][0]["paper_id"] = "paper-fixture-wrong"
    elif mutation == "nested_paper":
        validated_payload["citations"]["cite-alpha"]["evidence"][0]["paper_id"] = "paper-fixture-beta"
    elif mutation == "row_id":
        validated_payload["citations"]["cite-alpha"]["citation_id"] = "cite-beta"
    else:
        validated_payload["references"][0]["id"] = "paper-fixture-beta"
    with pytest.raises(JournalError) as error:
        writing_export.render_writing_docx(validated_payload)
    assert error.value.code in ("MISSING_REFERENCE", "CITATION_MISMATCH")


@pytest.mark.parametrize("mutation", ["outline_title", "outline_level", "explicit_section", "explicit_sections"])
def test_conflicting_draft_chapter_binding_rejected(validated_payload, mutation):
    if mutation == "outline_title":
        validated_payload["draft"]["outline"][0]["title"] = "Different draft title"
    elif mutation == "outline_level":
        validated_payload["draft"]["outline"][0]["level"] = 2
    elif mutation == "explicit_section":
        validated_payload["citations"]["cite-alpha"]["section_id"] = "results"
    else:
        validated_payload["citations"]["cite-alpha"]["section_ids"] = ["results"]
    with pytest.raises(JournalError) as error:
        writing_export.render_writing_docx(validated_payload)
    assert error.value.code == "CITATION_MISMATCH"


@pytest.mark.parametrize("field,value", [("reading_id", "reading-other"), ("file_sha256", "c" * 64), ("citation_id", "cite-beta")])
def test_nested_evidence_identity_rejected(validated_payload, field, value):
    validated_payload["citations"]["cite-alpha"]["evidence"][0][field] = value
    with pytest.raises(JournalError):
        writing_export.render_writing_docx(validated_payload)


@pytest.mark.parametrize("kind", ["agent_inference", "unverified"])
def test_inference_or_unverified_claim_cannot_masquerade_as_author_summary(validated_payload, kind):
    validated_payload["citations"]["cite-alpha"]["kind"] = kind
    with pytest.raises(JournalError):
        writing_export.render_writing_docx(validated_payload)


@pytest.mark.parametrize("field", ["citation_ids", "own_material_ids"])
def test_real_sources_are_required_by_paragraph_kind(validated_payload, field):
    paragraphs = validated_payload["draft"]["sections"][0]["paragraphs"]
    if field == "citation_ids":
        paragraphs[0][field] = []
    else:
        paragraphs[3][field] = ["UM-missing"]
    with pytest.raises(JournalError):
        writing_export.render_writing_docx(validated_payload)


@pytest.mark.parametrize("path,value", [
    (("revision",), True), (("profile",), []), (("profile", "language"), "fr"),
    (("draft", "language"), "en"), (("draft", "sections"), {}), (("references",), {}),
    (("references", 0, "title"), ""), (("references", 0, "authors"), "not-a-list"),
    (("references", 0, "year"), True), (("citations", "cite-alpha", "file_sha256"), "not-sha"),
    (("citations", "cite-alpha", "claim_index"), -1), (("citations", "cite-alpha", "kind"), []),
    (("citations", "cite-alpha", "evidence"), []),
    (("citations", "cite-alpha", "evidence", 0, "page_number"), 0),
    (("citations", "cite-alpha", "evidence", 0, "page_number"), True),
    (("citations", "cite-alpha", "evidence", 0, "quote"), ""),
    (("draft", "sections", 0, "paragraphs", 0, "text"), "bad\x00text"),
    (("draft", "sections", 0, "paragraphs", 0, "kind"), {}),
    (("draft", "sections", 0, "paragraphs", 0, "citation_ids"), ["cite-alpha", "cite-alpha"]),
    (("coverage",), []), (("gaps",), [123]),
])
def test_malformed_payload_is_rejected_as_domain_error(validated_payload, path, value):
    current = validated_payload
    for key in path[:-1]:
        current = current[key]
    current[path[-1]] = value
    with pytest.raises(JournalError):
        writing_export.render_writing_docx(validated_payload)


@pytest.mark.parametrize("payload", [None, [], {}, {"project_id": "fixture", "revision": 1}])
def test_missing_payload_fields_are_rejected(payload):
    with pytest.raises(JournalError):
        writing_export.render_writing_docx(payload)


def test_payload_budget_and_non_json_values_rejected(validated_payload):
    validated_payload["extra_fixture"] = "x" * (2 * 1024 * 1024)
    with pytest.raises(JournalError):
        writing_export.render_writing_docx(validated_payload)
    validated_payload["extra_fixture"] = object()
    with pytest.raises(JournalError):
        writing_export.render_writing_docx(validated_payload)


@pytest.mark.parametrize("url", ["file:///private/fixture.pdf", "https://user:FIXTURE-password@example.org/paper",
                                "https://example.org/paper?access_token=FIXTURE-secret", "https://example.org/paper?api_key=FIXTURE-secret",
                                "https://example.org/paper#access_token=FIXTURE-secret"])
def test_private_or_credentialed_reference_urls_rejected(validated_payload, url):
    validated_payload["references"][0]["landing_url"] = url
    with pytest.raises(JournalError) as error:
        writing_export.render_writing_docx(validated_payload)
    assert "FIXTURE-secret" not in str(error.value) and "FIXTURE-password" not in str(error.value)


def test_export_publishes_valid_docx_in_existing_parent(validated_payload, tmp_path):
    target = tmp_path / "中文 fixture Ω.docx"
    result = writing_export.export_writing_docx(validated_payload, target)
    assert result == target and result.is_file()
    assert "中文与 Ω" in all_text(Document(result))
    assert list(tmp_path.iterdir()) == [target]


def test_default_export_preserves_existing_file_without_rendering(validated_payload, tmp_path, monkeypatch):
    target = tmp_path / "fixture.docx"
    target.write_bytes(b"old fixture")
    monkeypatch.setattr(writing_export, "render_writing_docx", lambda payload: pytest.fail("must refuse existing target before rendering"))
    with pytest.raises(JournalError) as error:
        writing_export.export_writing_docx(validated_payload, str(target))
    assert error.value.code == "OUTPUT_EXISTS"
    assert target.read_bytes() == b"old fixture" and list(tmp_path.iterdir()) == [target]


def test_explicit_overwrite_replaces_atomically(validated_payload, tmp_path):
    target = tmp_path / "fixture.DOCX"
    target.write_bytes(b"old fixture")
    assert writing_export.export_writing_docx(validated_payload, str(target), overwrite=True) == target
    assert target.read_bytes().startswith(b"PK") and list(tmp_path.iterdir()) == [target]


@pytest.mark.parametrize("overwrite", [1, None, "true"])
def test_overwrite_requires_boolean(validated_payload, tmp_path, overwrite):
    with pytest.raises(JournalError):
        writing_export.export_writing_docx(validated_payload, tmp_path / "fixture.docx", overwrite=overwrite)
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize("path", ["relative.docx", "", "https://example.org/fixture.docx", "file:///tmp/fixture.docx",
                                 r"\\server\share\fixture.docx", "//server/share/fixture.docx", r"/\server\share\fixture.docx",
                                 r"\/server/share/fixture.docx", "bad\x00.docx", "bad\t.docx", None, 123])
def test_bad_output_paths_rejected(validated_payload, path):
    with pytest.raises(JournalError):
        writing_export.export_writing_docx(validated_payload, path)


@pytest.mark.parametrize("shape", ["wrong_suffix", "missing_parent", "directory_target", "dotdot"])
def test_non_docx_missing_parent_and_directory_target_rejected(validated_payload, tmp_path, shape):
    if shape == "wrong_suffix":
        target = tmp_path / "fixture.pdf"
    elif shape == "missing_parent":
        target = tmp_path / "nonexistent" / "fixture.docx"
    elif shape == "directory_target":
        target = tmp_path / "folder.docx"
        target.mkdir()
    else:
        target = tmp_path / ".." / "fixture.docx"
    before = set(tmp_path.iterdir())
    with pytest.raises(JournalError):
        writing_export.export_writing_docx(validated_payload, target)
    assert set(tmp_path.iterdir()) == before


@pytest.mark.parametrize("position", ["target", "ancestor"])
def test_symlink_output_or_parent_rejected(validated_payload, tmp_path, position):
    real = tmp_path / "real"
    real.mkdir()
    if position == "target":
        outside = real / "original.docx"
        outside.write_bytes(b"original fixture")
        link = tmp_path / "link.docx"
        try:
            link.symlink_to(outside)
        except OSError:
            pytest.skip("symlink creation privilege is unavailable")
        target = link
    else:
        link = tmp_path / "alias"
        try:
            link.symlink_to(real, target_is_directory=True)
        except OSError:
            pytest.skip("symlink creation privilege is unavailable")
        target = link / "fixture.docx"
    with pytest.raises(JournalError):
        writing_export.export_writing_docx(validated_payload, target, overwrite=True)
    assert not (real / "fixture.docx").exists()
    if position == "target":
        assert outside.read_bytes() == b"original fixture"


def test_junction_or_reparse_parent_is_rejected(validated_payload, tmp_path, monkeypatch):
    original_lstat = Path.lstat

    def reparse_lstat(path, *args, **kwargs):
        result = original_lstat(path, *args, **kwargs)
        if path == tmp_path:
            return SimpleNamespace(st_mode=stat.S_IFDIR, st_file_attributes=0x400)
        return result

    monkeypatch.setattr(Path, "lstat", reparse_lstat)
    with pytest.raises(JournalError):
        writing_export.export_writing_docx(validated_payload, tmp_path / "fixture.docx")
    assert not list(tmp_path.iterdir())


def test_invalid_payload_creates_no_staging_file(validated_payload, tmp_path):
    validated_payload["references"] = []
    with pytest.raises(JournalError):
        writing_export.export_writing_docx(validated_payload, tmp_path / "fixture.docx")
    assert not list(tmp_path.iterdir())


def test_racing_target_is_never_overwritten_and_staging_file_is_cleaned(validated_payload, tmp_path, monkeypatch):
    original_link = os.link
    target = tmp_path / "fixture.docx"

    def raced_link(source, destination):
        Path(destination).write_bytes(b"raced fixture")
        return original_link(source, destination)

    monkeypatch.setattr(os, "link", raced_link)
    with pytest.raises(JournalError) as error:
        writing_export.export_writing_docx(validated_payload, target)
    assert error.value.code == "OUTPUT_EXISTS"
    assert target.read_bytes() == b"raced fixture" and list(tmp_path.iterdir()) == [target]


@pytest.mark.parametrize("overwrite", [False, True])
def test_publish_failure_keeps_old_target_and_cleans_temporary_file(validated_payload, tmp_path, monkeypatch, overwrite):
    target = tmp_path / "fixture.docx"
    if overwrite:
        target.write_bytes(b"old fixture")

    def failed_publish(*args, **kwargs):
        raise OSError("fixture private C:\\secret-fixture\\paper.pdf token=DO-NOT-LEAK")

    monkeypatch.setattr(os, "replace" if overwrite else "link", failed_publish)
    with pytest.raises(JournalError) as error:
        writing_export.export_writing_docx(validated_payload, target, overwrite=overwrite)
    assert error.value.code == "EXPORT_FAILED"
    assert "DO-NOT-LEAK" not in str(error.value) and "secret-fixture" not in str(error.value)
    assert error.value.__suppress_context__ is True
    if overwrite:
        assert target.read_bytes() == b"old fixture" and list(tmp_path.iterdir()) == [target]
    else:
        assert not list(tmp_path.iterdir())


def test_failed_fsync_cleans_staging_file(validated_payload, tmp_path, monkeypatch):
    def failed_fsync(*args, **kwargs):
        raise OSError("fixture write failure")

    monkeypatch.setattr(os, "fsync", failed_fsync)
    with pytest.raises(JournalError):
        writing_export.export_writing_docx(validated_payload, tmp_path / "fixture.docx")
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize("field,value", [("paper_id", "paper-fixture-beta"), ("file_sha256", "c" * 64)])
def test_reading_identity_cannot_be_reused_for_another_paper_or_version(validated_payload, field, value):
    validated_payload["citations"]["cite-alpha-again"][field] = value
    with pytest.raises(JournalError) as error:
        writing_export.render_writing_docx(validated_payload)
    assert error.value.code == "CITATION_MISMATCH"


def test_long_title_is_preserved_in_body_without_core_property_overflow(validated_payload):
    title = "Fixture 中文长标题" * 40
    validated_payload["draft"]["title"] = title
    document, _ = rendered(validated_payload)
    assert document.paragraphs[0].text == title
    assert document.core_properties.title == title[:255]


def test_reparse_ancestor_is_checked_before_descendant(validated_payload, tmp_path, monkeypatch):
    target = tmp_path / "fixture.docx"
    original_lstat = Path.lstat

    def guarded_lstat(path, *args, **kwargs):
        if path == target:
            pytest.fail("must not inspect a child under a redirected parent")
        if path == tmp_path:
            return SimpleNamespace(st_mode=stat.S_IFDIR, st_file_attributes=0x400)
        return original_lstat(path, *args, **kwargs)

    monkeypatch.setattr(Path, "lstat", guarded_lstat)
    with pytest.raises(JournalError):
        writing_export.export_writing_docx(validated_payload, target)
    assert not list(tmp_path.iterdir())


def test_parent_redirected_during_render_is_rejected_before_staging(validated_payload, tmp_path, monkeypatch):
    original_render = writing_export.render_writing_docx
    original_lstat = Path.lstat
    redirected = False

    def redirect_after_render(payload):
        nonlocal redirected
        blob = original_render(payload)
        redirected = True
        return blob

    def redirected_lstat(path, *args, **kwargs):
        if redirected and path == tmp_path:
            return SimpleNamespace(st_mode=stat.S_IFDIR, st_file_attributes=0x400)
        return original_lstat(path, *args, **kwargs)

    monkeypatch.setattr(writing_export, "render_writing_docx", redirect_after_render)
    monkeypatch.setattr(Path, "lstat", redirected_lstat)
    monkeypatch.setattr(tempfile, "mkstemp", lambda **kwargs: pytest.fail("must not allocate a file under a redirected parent"))
    with pytest.raises(JournalError):
        writing_export.export_writing_docx(validated_payload, tmp_path / "fixture.docx")
    assert not list(tmp_path.iterdir())


def test_failed_fdopen_closes_descriptor_and_cleans_staging(validated_payload, tmp_path, monkeypatch):
    original_close = os.close
    closed = []

    def failed_fdopen(*args, **kwargs):
        raise OSError("fixture fdopen failure")

    def tracked_close(descriptor):
        closed.append(descriptor)
        return original_close(descriptor)

    monkeypatch.setattr(os, "fdopen", failed_fdopen)
    monkeypatch.setattr(os, "close", tracked_close)
    with pytest.raises(JournalError) as error:
        writing_export.export_writing_docx(validated_payload, tmp_path / "fixture.docx")
    assert error.value.code == "EXPORT_FAILED"
    assert closed and not list(tmp_path.iterdir())
