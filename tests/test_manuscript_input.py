import os
import sys
import tempfile
import zipfile
from unittest.mock import MagicMock, patch

import pytest
from paperflow.engine.journals.manuscript import prepare_manuscript
from paperflow.engine.journals.models import JournalError


def test_mutually_exclusive_and_empty():
    with pytest.raises(JournalError) as exc:
        prepare_manuscript()
    assert exc.value.code == "MUTUALLY_EXCLUSIVE_INPUT"

    with pytest.raises(JournalError) as exc:
        prepare_manuscript(text="hello", file_path="sample.txt")
    assert exc.value.code == "MUTUALLY_EXCLUSIVE_INPUT"

    with pytest.raises(JournalError) as exc:
        prepare_manuscript(text="   ")
    assert exc.value.code == "MUTUALLY_EXCLUSIVE_INPUT"


def test_invalid_max_chars_and_mode():
    with pytest.raises(JournalError) as exc:
        prepare_manuscript(text="hello world", max_chars=500)
    assert exc.value.code == "INVALID_MAX_CHARS"

    with pytest.raises(JournalError) as exc:
        prepare_manuscript(text="hello world", max_chars=200000)
    assert exc.value.code == "INVALID_MAX_CHARS"

    with pytest.raises(JournalError) as exc:
        prepare_manuscript(text="hello world", mode="invalid_mode")
    assert exc.value.code == "INVALID_MODE"


def test_direct_text_input_hash_stability():
    text1 = "This is a research proposal about machine learning."
    res1 = prepare_manuscript(text=text1, mode="auto")
    data1 = res1["data"]
    assert res1["status"] == "success"
    assert data1["mode"] == "idea"  # auto on text -> idea
    assert data1["source_type"] == "text"
    assert data1["source_name"] == ""
    assert data1["text"] == text1
    assert data1["truncated"] is False
    assert res1["coverage"]["text_extraction_only"] is True
    assert any("严禁直接执行" in opt for opt in res1["suggested_options"])

    # 相同的 text 和 mode，input_id 必须相同
    res2 = prepare_manuscript(text=text1, mode="auto")
    assert res2["data"]["input_id"] == data1["input_id"]

    # 改变 mode，input_id 发生改变
    res_manuscript = prepare_manuscript(text=text1, mode="manuscript")
    assert res_manuscript["data"]["mode"] == "manuscript"
    assert res_manuscript["data"]["input_id"] != data1["input_id"]

    # 改变 text，input_id 发生改变
    res3 = prepare_manuscript(text=text1 + " Extra.", mode="auto")
    assert res3["data"]["input_id"] != data1["input_id"]


def test_txt_and_md_utf8_sig_and_truncation():
    with tempfile.TemporaryDirectory() as tmpdir:
        # TXT 带 BOM
        txt_path = os.path.join(tmpdir, "paper_sample.txt")
        long_content = "第一段：研究背景介绍。\n\n第二段：相关工作梳理与对比。" * 100
        with open(txt_path, "wb") as f:
            f.write(long_content.encode("utf-8-sig"))

        res = prepare_manuscript(file_path=txt_path, mode="auto", max_chars=1200)
        data = res["data"]
        assert data["mode"] == "manuscript"  # auto on file -> manuscript
        assert data["source_type"] == "txt"
        assert data["source_name"] == "paper_sample.txt"
        assert data["truncated"] is True
        assert data["text_length"] == 1200
        assert len(data["text"]) == 1200
        assert len(res["warnings"]) > 0

        # 换另一个目录，但相同文件名和内容，input_id 必须完全一致（不包含绝对路径）
        with tempfile.TemporaryDirectory() as tmpdir2:
            txt_path2 = os.path.join(tmpdir2, "paper_sample.txt")
            with open(txt_path2, "wb") as f2:
                f2.write(long_content.encode("utf-8-sig"))
            res2 = prepare_manuscript(file_path=txt_path2, mode="auto", max_chars=1200)
            assert res2["data"]["input_id"] == data["input_id"]

        # MD 文件
        md_path = os.path.join(tmpdir, "notes.md")
        with open(md_path, "w", encoding="utf-8") as f:
            f.write("# 标题\n\n正文内容")
        res_md = prepare_manuscript(file_path=md_path)
        assert res_md["data"]["source_type"] == "md"
        assert res_md["data"]["source_name"] == "notes.md"
        assert "# 标题" in res_md["data"]["text"]


def test_illegal_encoding():
    with tempfile.TemporaryDirectory() as tmpdir:
        txt_path = os.path.join(tmpdir, "gbk_file.txt")
        with open(txt_path, "wb") as f:
            f.write("中文测试GBK编码".encode("gbk"))

        with pytest.raises(JournalError) as exc:
            prepare_manuscript(file_path=txt_path)
        assert exc.value.code == "ILLEGAL_ENCODING"


def test_network_and_unc_paths():
    with pytest.raises(JournalError) as exc:
        prepare_manuscript(file_path=r"\\192.168.1.1\share\paper.docx")
    assert exc.value.code == "NETWORK_PATH_REJECTED"

    with pytest.raises(JournalError) as exc:
        prepare_manuscript(file_path="http://example.com/paper.pdf")
    assert exc.value.code == "NETWORK_PATH_REJECTED"


def test_docx_with_table_and_immutable():
    import docx

    with tempfile.TemporaryDirectory() as tmpdir:
        docx_path = os.path.join(tmpdir, "test_doc.docx")
        doc = docx.Document()
        doc.add_heading("论文标题", level=1)
        doc.add_paragraph("这是正文第一段。")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "指标"
        table.cell(0, 1).text = "数值"
        table.cell(1, 0).text = "准确率"
        table.cell(1, 1).text = "95.5%"
        doc.save(docx_path)

        # 记录原始文件大小和修改时间
        mtime_before = os.path.getmtime(docx_path)
        size_before = os.path.getsize(docx_path)

        res = prepare_manuscript(file_path=docx_path)
        data = res["data"]
        assert data["source_type"] == "docx"
        assert "论文标题" in data["text"]
        assert "这是正文第一段。" in data["text"]
        assert "指标 | 数值" in data["text"]
        assert "准确率 | 95.5%" in data["text"]

        # 检查原文件未被修改
        assert os.path.getmtime(docx_path) == mtime_before
        assert os.path.getsize(docx_path) == size_before


def test_docx_zip_limits():
    with tempfile.TemporaryDirectory() as tmpdir:
        # 使用真实压缩的高冗余零字节，使得压缩文件体积 < 32 MiB，但解压体积 > 32 MiB
        fake_docx = os.path.join(tmpdir, "bomb.docx")
        with zipfile.ZipFile(fake_docx, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("word/document.xml", b"0" * (33 * 1024 * 1024))

        with pytest.raises(JournalError) as exc:
            prepare_manuscript(file_path=fake_docx)
        assert exc.value.code == "ZIP_BOMB_DETECTED"

        # 路径遍历成员模拟
        bad_path_docx = os.path.join(tmpdir, "bad_path.docx")
        with zipfile.ZipFile(bad_path_docx, "w") as zf:
            zf.writestr("../outside.txt", b"malicious")

        with pytest.raises(JournalError) as exc:
            prepare_manuscript(file_path=bad_path_docx)
        assert exc.value.code == "DANGEROUS_ARCHIVE"


def test_pdf_dependency_missing():
    # 当环境中无 pypdf 或 mock ImportError
    with tempfile.TemporaryDirectory() as tmpdir:
        pdf_path = os.path.join(tmpdir, "sample.pdf")
        with open(pdf_path, "wb") as f:
            f.write(b"%PDF-1.4 dummy content")

        with patch.dict(sys.modules, {"pypdf": None}):
            with pytest.raises(JournalError) as exc:
                prepare_manuscript(file_path=pdf_path)
            assert exc.value.code == "DEPENDENCY_MISSING"
            assert "paperflow-mcp[pdf]" in str(exc.value)
            assert "paperflow[pdf]" not in str(exc.value)


def test_input_type_and_size_limits_and_unsupported_ext():
    # 非字符串类型校验（即使提供了合法 file_path 也不能静默忽略）
    with tempfile.TemporaryDirectory() as tmpdir:
        txt_path = os.path.join(tmpdir, "valid.txt")
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write("valid content")

        with pytest.raises(JournalError) as exc:
            prepare_manuscript(text=123, file_path=txt_path)  # type: ignore
        assert exc.value.code == "INVALID_INPUT_TYPE"

        with pytest.raises(JournalError) as exc:
            prepare_manuscript(text="valid", file_path=None)  # type: ignore
        assert exc.value.code == "INVALID_INPUT_TYPE"

        # 未知扩展名不回显后缀
        secret_ext_path = os.path.join(tmpdir, "file.secret_custom_ext")
        with open(secret_ext_path, "w", encoding="utf-8") as f:
            f.write("some text")

        with pytest.raises(JournalError) as exc:
            prepare_manuscript(file_path=secret_ext_path)
        assert exc.value.code == "UNSUPPORTED_FILE_TYPE"
        assert "secret_custom_ext" not in str(exc.value)

    # 直接文本超过 32 MiB 限制
    huge_text = "A" * (32 * 1024 * 1024 + 10)
    with pytest.raises(JournalError) as exc:
        prepare_manuscript(text=huge_text)
    assert exc.value.code == "TEXT_TOO_LARGE"


def test_pdf_multi_page_truncation_and_early_stop():
    with tempfile.TemporaryDirectory() as tmpdir:
        pdf_path = os.path.join(tmpdir, "many_pages.pdf")
        with open(pdf_path, "wb") as f:
            f.write(b"%PDF-1.4 dummy content")

        # 场景1：120 页，每页只有短文本（总字符数远小于 max_chars），但因超过 100 页仍须标记 truncated=True
        mock_reader_120 = MagicMock()
        mock_reader_120.is_encrypted = False
        pages_120 = []
        for i in range(120):
            p = MagicMock()
            p.extract_text.return_value = f"Page {i + 1} short."
            pages_120.append(p)
        mock_reader_120.pages = pages_120

        with patch("pypdf.PdfReader", return_value=mock_reader_120):
            res = prepare_manuscript(file_path=pdf_path, max_chars=60000)
            data = res["data"]
            assert data["truncated"] is True
            assert data["pages_read"] == 100
            assert data["total_pages"] == 120
            # 第 101 页不应被调用
            assert pages_120[100].extract_text.call_count == 0

        # 场景2：10 页，每页 600 字符，设置 max_chars=1500，应在读完第 3 页达到字符预算后立即停止后续页提取
        mock_reader_budget = MagicMock()
        mock_reader_budget.is_encrypted = False
        pages_budget = []
        for i in range(10):
            p = MagicMock()
            p.extract_text.return_value = "X" * 600
            pages_budget.append(p)
        mock_reader_budget.pages = pages_budget

        with patch("pypdf.PdfReader", return_value=mock_reader_budget):
            res_b = prepare_manuscript(file_path=pdf_path, max_chars=1500)
            data_b = res_b["data"]
            assert data_b["truncated"] is True
            assert data_b["pages_read"] == 3
            assert data_b["total_pages"] == 10
            assert data_b["text_length"] == 1500
            # 第 4 页及之后不应再提取
            assert pages_budget[3].extract_text.call_count == 0


def test_pdf_scanned_ocr_required_and_encrypted():
    with tempfile.TemporaryDirectory() as tmpdir:
        pdf_path = os.path.join(tmpdir, "scanned.pdf")
        with open(pdf_path, "wb") as f:
            f.write(b"%PDF-1.4 dummy content")

        mock_pypdf = MagicMock()
        mock_reader = MagicMock()
        mock_reader.is_encrypted = False
        mock_page = MagicMock()
        mock_page.get_contents.return_value = None
        mock_page.extract_text.return_value = "   "  # 无可读文字
        mock_reader.pages = [mock_page]
        mock_pypdf.PdfReader.return_value = mock_reader

        with patch.dict(sys.modules, {"pypdf": mock_pypdf}):
            with pytest.raises(JournalError) as exc:
                prepare_manuscript(file_path=pdf_path)
            assert exc.value.code == "OCR_REQUIRED"

        # 加密测试
        mock_encrypted_reader = MagicMock()
        mock_encrypted_reader.is_encrypted = True
        mock_encrypted_reader.decrypt.return_value = 0
        mock_pypdf.PdfReader.return_value = mock_encrypted_reader

        with patch.dict(sys.modules, {"pypdf": mock_pypdf}):
            with pytest.raises(JournalError) as exc:
                prepare_manuscript(file_path=pdf_path)
            assert exc.value.code == "PDF_ENCRYPTED"


def test_pdf_decompression_resource_limit():
    import pypdf
    from pypdf.errors import LimitReachedError

    with tempfile.TemporaryDirectory() as tmpdir:
        pdf_path = os.path.join(tmpdir, "heavy.pdf")
        with open(pdf_path, "wb") as f:
            f.write(b"%PDF-1.4 dummy content")

        mock_reader = MagicMock()
        mock_reader.is_encrypted = False
        mock_page = MagicMock()
        # 模拟解码超限抛出 pypdf 自身的 LimitReachedError
        mock_page.extract_text.side_effect = LimitReachedError("Limit reached while decompressing.")
        mock_reader.pages = [mock_page]

        with patch("pypdf.PdfReader", return_value=mock_reader):
            with pytest.raises(JournalError) as exc:
                prepare_manuscript(file_path=pdf_path)
            assert exc.value.code == "PDF_RESOURCE_LIMIT"


def test_real_pdf_text_extraction():
    import pypdf

    # 使用真实 pypdf 创建并提取一个合法的 PDF
    writer = pypdf.PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    # 纯空白页面当做纯扫描或无文本
    with tempfile.TemporaryDirectory() as tmpdir:
        blank_pdf = os.path.join(tmpdir, "blank.pdf")
        with open(blank_pdf, "wb") as f:
            writer.write(f)

        with pytest.raises(JournalError) as exc:
            prepare_manuscript(file_path=blank_pdf)
        assert exc.value.code == "OCR_REQUIRED"
