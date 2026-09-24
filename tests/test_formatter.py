"""Unit tests for PaperFormatAuditor and PaperFormatNormalizer."""

import os
import tempfile
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from paperflow.engine.formatter import IssueSeverity, PaperFormatAuditor, PaperFormatNormalizer


class TestPaperFormatEngine(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.raw_docx = os.path.join(self.tmp.name, "raw_paper.docx")
        self.fixed_docx = os.path.join(self.tmp.name, "fixed_paper.docx")

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def test_punctuation_normalization(self) -> None:
        raw_text = "随着深度学习模型的演进,计算复杂度急剧攀升;在实时场景(如无人驾驶)中,延迟高达 3.14ms [1,2],难以满足需求."
        normalized = PaperFormatNormalizer.normalize_chinese_punctuation(raw_text)
        # 半角逗号、分号、冒号、括号应被转换为全角
        self.assertIn("演进，", normalized)
        self.assertIn("攀升；", normalized)
        self.assertIn("（如无人驾驶）", normalized)
        self.assertIn("中，", normalized)
        # 小数点 3.14 和方括号引用 [1,2] 应完好保留，不被误篡改
        self.assertIn("3.14ms", normalized)
        self.assertIn("[1,2]", normalized)

    def test_audit_detects_defects(self) -> None:
        doc = Document()
        # 1. 论文大标题误用 Heading 1
        doc.add_heading("基于边缘协同的轻量级目标检测模型研究", level=1)
        # 2. 连续空行
        doc.add_paragraph("")
        doc.add_paragraph("")
        # 3. 伪标题（普通段落手写 '第1章 绪论'）
        doc.add_paragraph("第1章 绪论")
        # 4. 正文缺失首行缩进且存在半角标点
        doc.add_paragraph("在工业物联网边缘设备中,传感器流式数据的高通量特征与算力受限构成核心矛盾.")
        # 5. 跨级标题：直接跳到 Heading 3
        doc.add_heading("1.1.1 模型轻量化机制", level=3)
        # 6. 带竖线的表格
        table = doc.add_table(rows=2, cols=2)
        tbl_pr = table._tbl.tblPr
        tbl_borders = OxmlElement("w:tblBorders")
        inside_v = OxmlElement("w:insideV")
        inside_v.set(qn("w:val"), "single")
        tbl_borders.append(inside_v)
        tbl_pr.append(tbl_borders)

        # 7. 断号引用 [1] 然后 [3]（缺少 [2]）
        doc.add_paragraph("前人实验结果表明算法性能突出 [1]，后续模型也延续了此方法 [3]。")

        doc.save(self.raw_docx)

        # 运行审计
        report = PaperFormatAuditor.audit_docx(self.raw_docx)
        self.assertEqual(report["status"], "success")
        self.assertFalse(report["is_compliant"])
        self.assertLess(report["format_score"], 90)

        rule_ids = [issue["rule_id"] for issue in report["issues"]]
        self.assertIn("TITLE_OUTLINE_LEAK", rule_ids)
        self.assertIn("CONSECUTIVE_BLANK_LINES", rule_ids)
        self.assertIn("PSEUDO_HEADING", rule_ids)
        self.assertIn("BODY_INDENT_MISSING", rule_ids)
        self.assertIn("PUNCTUATION_HALF_WIDTH", rule_ids)
        self.assertIn("HEADING_LEVEL_SKIPPED", rule_ids)
        self.assertIn("NON_THREE_LINE_TABLE", rule_ids)
        self.assertIn("CITATION_DISCONTINUITY", rule_ids)

    def test_normalize_docx_auto_heals(self) -> None:
        doc = Document()
        # 标题和多余空行
        doc.add_paragraph("基于边缘协同的轻量级目标检测模型研究")
        doc.add_paragraph("")
        doc.add_paragraph("")
        doc.add_paragraph("")
        # 正文（手打空格+半角标点）
        doc.add_paragraph("　　在边缘计算落地过程中,传统网络存在海量参数冗余,计算开销巨大.")
        # 普通网格表格
        table = doc.add_table(rows=3, cols=3)
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                cell.text = f"Data-{r_idx}-{c_idx}"

        doc.save(self.raw_docx)

        # 运行一键规范化自愈
        res = PaperFormatNormalizer.normalize_docx(self.raw_docx, self.fixed_docx)
        self.assertEqual(res["status"], "success")
        self.assertTrue(os.path.isfile(self.fixed_docx))
        self.assertGreater(res["total_changes"], 0)

        # 重新审查修复后的文件，排版应得到明显纠正
        fixed_doc = Document(self.fixed_docx)
        # 连续空行应被清理
        texts = [p.text for p in fixed_doc.paragraphs]
        blank_streak = 0
        for t in texts:
            if not t.strip():
                blank_streak += 1
                self.assertLess(blank_streak, 2, "连续空行未被彻底清理")
            else:
                blank_streak = 0

        # 正文标点应已转换为全角
        body_text = "".join(texts)
        self.assertIn("过程中，传统网络", body_text)

        # 表格应已被转为学术三线表
        fixed_report = PaperFormatAuditor.audit_docx(self.fixed_docx)
        fixed_rules = [i["rule_id"] for i in fixed_report["issues"]]
        self.assertNotIn("NON_THREE_LINE_TABLE", fixed_rules)
        self.assertNotIn("CONSECUTIVE_BLANK_LINES", fixed_rules)

    def test_reviewer_first_impression_audit_rules(self) -> None:
        """测试审稿人第一印象扣分项检测：空泛标题、缩写无首定义、悬空标题、摘要结论重复。"""
        doc = Document()
        # 1. 题目假大空（缺少具体机理/物理量/核心矛盾）
        doc.add_paragraph("基于深度学习的图像识别系统研究")

        # 2. 摘要与结论高度重叠
        abs_text = "本文针对目标检测任务展开研究，构建了深度卷积网络架构并完成了模型训练。实验表明该模型在测试集上达到了92.5%的平均精度，展现了良好的分类效果。"
        doc.add_paragraph("摘要")
        doc.add_paragraph(abs_text)

        # 3. 悬空标题：Heading 1 紧跟 Heading 2，缺少承上启下过渡段落
        doc.add_heading("1 算法体系与结构设计", level=1)
        doc.add_heading("1.1 特征提取模块", level=2)

        # 4. 正文中英文大写缩写首次出现无完整英文全称定义（如 CFD, CBAM）
        doc.add_paragraph("在特征提取阶段，本研究采用 CFD 模拟流场数据，并引入 CBAM 机制以提高显著性表征。符合国家标准 GB 和 ISO 规范。")

        # 5. 结论段落直接机械复制摘要句子
        doc.add_heading("2 结论", level=1)
        conc_text = "本文针对目标检测任务展开研究，构建了深度卷积网络架构并完成了模型训练。实验表明该模型在测试集上达到了92.5%的平均精度，展现了良好的分类效果。"
        doc.add_paragraph(conc_text)

        doc.save(self.raw_docx)

        report = PaperFormatAuditor.audit_docx(self.raw_docx)
        self.assertEqual(report["status"], "success")

        issues_by_rule = {i["rule_id"]: i for i in report["issues"]}

        # 验证 4 个新增规则全部触发
        self.assertIn("TITLE_VAGUE_IMPRESSION", issues_by_rule)
        self.assertEqual(issues_by_rule["TITLE_VAGUE_IMPRESSION"]["severity"], "warning")
        self.assertIn("建议改为‘[具体机理/方法]+[核心矛盾]+[具体工程/物理对象]’形式", issues_by_rule["TITLE_VAGUE_IMPRESSION"]["message"])

        self.assertIn("ABBREVIATION_NO_DEFINITION", issues_by_rule)
        self.assertEqual(issues_by_rule["ABBREVIATION_NO_DEFINITION"]["severity"], "warning")
        self.assertIn("缩写首次出现无完整英文全称定义", issues_by_rule["ABBREVIATION_NO_DEFINITION"]["suggestion"])

        self.assertIn("DANGLING_HEADING", issues_by_rule)
        self.assertEqual(issues_by_rule["DANGLING_HEADING"]["severity"], "warning")
        self.assertIn("标题下方缺少承上启下过渡段落", issues_by_rule["DANGLING_HEADING"]["message"])

        self.assertIn("ABSTRACT_CONCLUSION_REPETITION", issues_by_rule)
        self.assertEqual(issues_by_rule["ABSTRACT_CONCLUSION_REPETITION"]["severity"], "warning")
        self.assertIn("摘要与结论文字高度重叠", issues_by_rule["ABSTRACT_CONCLUSION_REPETITION"]["message"])

    def test_reviewer_rules_clean_document_passes(self) -> None:
        """测试规范文档不会误报审稿人扣分规则。"""
        doc = Document()
        # 1. 具象标题（包含机理、物理量与工程对象）
        doc.add_paragraph("面向极端热冲击工况的多孔拓扑优化方法研究")

        # 2. 差异化摘要
        doc.add_paragraph("摘要：针对高超声速飞行器热防护结构热-力耦合退化难题，本文提出一种基于微观晶格形貌的自适应多孔拓扑优化方法。")

        # 3. 正常标题过渡：Heading 1 带有引言段落，再进入 Heading 2
        doc.add_heading("1 引言", level=1)
        doc.add_paragraph("高超声速热防护是航天领域的核心瓶颈。为阐明其内在热应力演化机制，本章首先介绍基础控制理论。")
        doc.add_heading("1.1 问题描述", level=2)

        # 4. 首次出现缩写附带完整英文全称：Full Name (ABBR)
        doc.add_paragraph("在数值仿真中，我们采用计算流体力学（Computational Fluid Dynamics, CFD）进行求解，并结合国家标准 GB 验证边界。")

        # 5. 差异化结论（定量成果与局限展望）
        doc.add_heading("4 结论与展望", level=1)
        doc.add_paragraph("本文系统揭示了微观孔隙拓扑对边界对流换热系数的调控机制。实测最高温度降低了14.2%，后续将继续验证超高马赫数环境。")

        doc.save(self.raw_docx)

        report = PaperFormatAuditor.audit_docx(self.raw_docx)
        self.assertEqual(report["status"], "success")
        rule_ids = [i["rule_id"] for i in report["issues"]]

        self.assertNotIn("TITLE_VAGUE_IMPRESSION", rule_ids)
        self.assertNotIn("ABBREVIATION_NO_DEFINITION", rule_ids)
        self.assertNotIn("DANGLING_HEADING", rule_ids)
        self.assertNotIn("ABSTRACT_CONCLUSION_REPETITION", rule_ids)


if __name__ == "__main__":
    unittest.main()
