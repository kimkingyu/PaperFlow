"""Offline writing contracts: real metadata, native judgement and rechecked evidence."""
import copy
import json
import socket
import sys
import types
from pathlib import Path

import pytest

from test_literature_service import FakeProviders, card_for, imported, pdf_bytes, record
from paperflow.engine.journals.models import JournalError
from paperflow.engine.literature.models import PaperRecord, paper_id_for
from paperflow.engine.literature.service import LiteratureService
from paperflow.engine.literature.writing_models import MAX_INPUT_BYTES, reliable_identity
from paperflow.engine.literature.writing_service import PaperWritingService


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError("writing core tests forbid real network")
    monkeypatch.setattr(socket, "create_connection", forbidden)
    from paperflow.engine.literature import providers
    monkeypatch.setattr(providers.http, "safe_fetch", forbidden)


@pytest.fixture
def writing(tmp_path):
    literature = LiteratureService(str(tmp_path / "library"), providers=FakeProviders())
    return PaperWritingService(literature=literature)


def profile(**kwargs):
    return {"title": "Evidence-backed literature draft", "research_question": "Which reproducible methods support deployment?", **kwargs}


def query(number=1, **kwargs):
    return {"id": "Q" + str(number), "query": "deployment " + str(number),
            "purpose": "Find reproducible methods for RQ1", "question_ids": ["RQ1"], **kwargs}


def assessment(candidate, **kwargs):
    return {"paper_id": candidate["paper_id"], "metadata_sha256": candidate["metadata_sha256"],
            "relevance": "core", "reason": "Its stated method directly addresses deployment constraints in RQ1.",
            "question_ids": ["RQ1"], "basis": "metadata", "limitations": ["Full text initially not read"], **kwargs}


def local_project(writing, tmp_path, count=1, kinds=None, own_materials=None):
    ids = []
    for number in range(count):
        file = tmp_path / ("fixture-" + str(number) + ".pdf")
        file.write_bytes(pdf_bytes(["A real reproducible method " + str(number) + ".", "Observed baseline " + str(number) + "."]))
        pid = writing.literature.import_pdf(str(file), "Method " + str(number))["data"]["paper_id"]
        reading = card_for(writing.literature.read(pid))
        if kinds:
            reading["claims"][0]["kind"] = kinds[number]
        writing.literature.save_reading(pid, reading)
        ids.append(pid)
    created = writing.create(profile(own_materials=own_materials or []), paper_ids=ids)["data"]
    return created


def selected_project(writing, tmp_path, count=1, kinds=None, own_materials=None):
    created = local_project(writing, tmp_path, count, kinds, own_materials)
    return writing.assess(created["project_id"], [assessment(c) for c in created["candidates"]], 1)["data"]


def draft_for(project, citation=None, kind="literature_summary", **kwargs):
    citation = citation or project["evidence_matrix"][0]["citation_id"]
    return {"title": "Reproducible methods: a literature-supported draft", "language": project["profile"]["language"],
            "outline": [{"id": "related_work", "title": "Related work", "level": 1,
                         "purpose": "Compare the stated methods", "evidence_ids": [citation]}],
            "sections": [{"id": "related_work", "title": "Related work", "level": 1,
                          "paragraphs": [{"text": "The referenced study describes a reproducible method.",
                                          "kind": kind, "citation_ids": [citation], "own_material_ids": []}]}], **kwargs}


def fake_export(monkeypatch, fail=None):
    module = types.ModuleType("paperflow.engine.literature.writing_export")
    payloads = []
    def render(payload):
        payloads.append(payload)
        if fail:
            raise fail
        return b"verified docx fixture"
    def export(payload, output_path, overwrite=False):
        payloads.append(payload)
        if fail:
            raise fail
        return Path(output_path)
    module.render_writing_docx = render
    module.export_writing_docx = export
    monkeypatch.setitem(sys.modules, module.__name__, module)
    return payloads


def test_construct_prepare_list_and_missing_get_are_inert(writing):
    assert not writing.library.directory.exists()
    prepared = writing.prepare("User supplied actual research question")
    assert set(prepared) == {"status", "error_code", "message", "data", "sources", "coverage", "warnings", "suggested_options"}
    assert prepared["data"]["input_id"].startswith("input-")
    assert all(k in prepared["data"] for k in ("profile_schema", "query_schema", "assessment_schema", "draft_schema", "agent_contract"))
    assert writing.list()["data"] == {"projects": [], "total": 0}
    with pytest.raises(JournalError):
        writing.get("writing-" + "a" * 24)
    assert not writing.library.directory.exists()
    assert not writing.literature.providers.calls


def test_reading_old_library_does_not_migrate(writing):
    writing.library.put(record())
    with writing.library.connection() as conn:
        before = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    assert writing.list()["data"]["total"] == 0
    with writing.library.connection() as conn:
        after = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    assert before == after
    assert "writing_projects" not in after


def test_default_question_and_candidates_are_not_fabricated_related(writing):
    paper = writing.library.put(record())
    created = writing.create(profile(sub_questions=[]), paper_ids=[paper["paper_id"]])["data"]
    assert created["revision"] == 1
    assert created["profile"]["sub_questions"] == [{"id": "RQ1", "question": profile()["research_question"]}]
    assert created["selected_paper_ids"] == [] and created["assessments"] == []
    assert created["stage"] == "needs_agent_assessment"
    assert created["evidence_matrix"] == []
    assert "relevance" not in created["candidates"][0]
    assert created["own_materials_verification"].startswith("caller-supplied")
    assert not writing.literature.providers.calls


@pytest.mark.parametrize("bad", [None, [], {}, "", "   ", True])
def test_invalid_prepare_is_sanitized_and_inert(writing, bad):
    with pytest.raises(JournalError) as exc:
        writing.prepare(bad)
    assert exc.value.code == "INVALID_INPUT"
    assert not writing.library.directory.exists()


@pytest.mark.parametrize("changes", [
    {"api_key": "NEVER-LEAK-SECRET"}, {"title": " "}, {"language": "fr"}, {"title": True},
    {"sub_questions": [{"id": "RQ1", "question": "x"}, {"id": "RQ1", "question": "y"}]},
    {"own_materials": [{"id": "UM1", "text": ""}]},
])
def test_profile_shape_forbids_extras_and_fake_material(writing, changes):
    with pytest.raises(JournalError) as exc:
        writing.create(profile(**changes))
    assert "NEVER-LEAK" not in str(exc.value)
    assert not writing.library.directory.exists()


def test_source_identity_and_own_materials_are_bound(writing):
    text = "Actual user material: measured throughput was 12 units."
    prepared = writing.prepare(text)["data"]
    with pytest.raises(JournalError) as exc:
        writing.create(profile(), source_text=text + " changed", input_id=prepared["input_id"])
    assert exc.value.code == "INPUT_MISMATCH"
    with pytest.raises(JournalError) as exc:
        writing.create(profile(), input_id=prepared["input_id"])
    assert exc.value.code == "INPUT_MISMATCH"
    with pytest.raises(JournalError) as exc:
        writing.create(profile(own_materials=[{"id": "UM1", "text": "invented experiments"}]),
                       source_text=text, input_id=prepared["input_id"])
    assert exc.value.code == "OWN_MATERIAL_MISMATCH"
    project = writing.create(profile(own_materials=[{"id": "UM1", "text": "measured throughput was 12 units."}]),
                             source_text=text, input_id=prepared["input_id"])["data"]
    prepared_writing = writing.prepare_writing(project["project_id"])
    assert prepared_writing["data"]["writing_materials"]["source_text"] == text
    assert not prepared_writing["coverage"]["own_materials_independently_verified"]


@pytest.mark.parametrize("changes", [
    {"year_from": True}, {"year_to": "2020"}, {"year_from": 2021, "year_to": 2020},
    {"question_ids": ["not-RQ"]}, {"sources": []}, {"sources": ["private"]}, {"query": ""},
    {"purpose": ""}, {"sources": ["crossref", "crossref"]}, {"api_key": "SECRET"},
])
def test_query_validation_precedes_network(writing, changes):
    with pytest.raises(JournalError):
        writing.create(profile(), queries=[query(**changes)])
    assert not writing.literature.providers.calls
    assert not writing.library.directory.exists()


def test_six_queries_ten_each_and_sixty_candidates(writing):
    class ManyProviders(FakeProviders):
        def search(self, source_id, text, **kwargs):
            self.calls.append((source_id, text, kwargs))
            number = text.split()[-1]
            return [record(source_id, number + "-" + str(i), doi="10.1234/" + number + source_id + str(i)) for i in range(12)]
    writing.literature.providers = ManyProviders()
    created = writing.create(profile(), queries=[query(i) for i in range(1, 7)])["data"]
    found = writing.search(created["project_id"], 1)["data"]
    assert len(found["candidates"]) == 60
    assert len(found["search_statuses"]) == 12
    for q in found["queries"]:
        assert sum(q["id"] in c["matched_query_ids"] for c in found["candidates"]) == 10
    assert all(call[2]["limit"] == 10 for call in writing.literature.providers.calls)
    assert all(s["count"] == 12 and s["returned_count"] == 10 and s["kept_count"] == 5 for s in found["search_statuses"])
    assert found["stage"] == "needs_agent_assessment" and found["assessments"] == []
    assert all("relevance_score" not in c and "score" not in c for c in found["candidates"])


@pytest.mark.parametrize("kwargs", [{"per_query_limit": True}, {"per_query_limit": 0}, {"per_query_limit": 11},
                                     {"queries": [query(i) for i in range(7)]}, {"queries": [query(), query()]}])
def test_search_limits_before_provider(writing, kwargs):
    created = writing.create(profile(), queries=[query()])["data"]
    with pytest.raises(JournalError):
        writing.search(created["project_id"], 1, **kwargs)
    assert not writing.literature.providers.calls
    assert writing.get(created["project_id"])["data"]["revision"] == 1


def test_partial_sources_keep_real_query_status_without_secret_errors(writing):
    writing.literature.providers.records["arxiv"] = JournalError("RATE_LIMITED", "token=DO-NOT-LEAK private traceback")
    created = writing.create(profile(), queries=[query(1), query(2)])["data"]
    out = writing.search(created["project_id"], 1)
    assert out["status"] == "partial"
    assert len(out["data"]["candidates"]) == 1
    assert out["data"]["candidates"][0]["matched_query_ids"] == ["Q1", "Q2"]
    assert len(out["data"]["search_statuses"]) == 4
    assert all(s["error_code"] == "RATE_LIMITED" for s in out["data"]["search_statuses"] if s["status"] == "error")
    assert "DO-NOT-LEAK" not in json.dumps(out)
    assert "traceback" not in json.dumps(out)


def test_all_sources_failure_is_partial_not_success_empty(writing):
    writing.literature.providers.records = {s: RuntimeError("APIKEY-SECRET") for s in ("crossref", "arxiv")}
    created = writing.create(profile(), queries=[query()])["data"]
    out = writing.search(created["project_id"], 1)
    assert out["status"] == "partial" and not out["data"]["candidates"]
    assert len(out["data"]["search_statuses"]) == 2
    assert all(s["status"] == "error" for s in out["data"]["search_statuses"])
    assert "APIKEY-SECRET" not in json.dumps(out)


def test_reliable_doi_merges_provenance_but_titles_and_versions_do_not(writing):
    a = record(identifier="doi-a", doi="https://doi.org/10.1234/SAME", publication_type="journal-article")
    b = record(identifier="doi-b", doi="10.1234/same")
    c = record(identifier="invalid-a", doi="not-a-doi")
    d = record(identifier="invalid-b", doi="not-a-doi")
    pre7 = record("arxiv", "1706.03762v7", doi="10.1234/same", arxiv_id="1706.03762v7", publication_type="preprint")
    pre8 = record("arxiv", "1706.03762v8", doi="10.1234/same", arxiv_id="1706.03762v8", publication_type="preprint")
    writing.literature.providers.records = {"crossref": [a, b, c, d], "arxiv": [pre7, pre8]}
    created = writing.create(profile(), queries=[query()])["data"]
    result = writing.search(created["project_id"], 1)["data"]
    assert len(result["candidates"]) == 5
    merged = next(candidate for candidate in result["candidates"] if len(candidate["source_records"]) == 2)
    assert {s["paper_id"] for s in merged["source_records"]} == {a["paper_id"], b["paper_id"]}
    assert len({c["paper"]["title"] for c in result["candidates"]}) == 1
    assert reliable_identity(pre7) != reliable_identity(pre8)


def test_assessment_batches_merge_and_native_reason_is_preserved(writing, tmp_path):
    created = local_project(writing, tmp_path, 2)
    first, second = created["candidates"]
    out = writing.assess(created["project_id"], [assessment(first)], 1, origin="nativeAgent")["data"]
    assert out["revision"] == 2 and len(out["assessments"]) == 1
    assert out["assessments"][0]["origin"] == "nativeAgent"
    out = writing.assess(created["project_id"], [assessment(second, relevance="background")], 2)["data"]
    assert len(out["assessments"]) == 2
    assert set(out["selected_paper_ids"]) == {first["paper_id"], second["paper_id"]}
    assert out["stage"] == "evidence_ready"
    out = writing.assess(created["project_id"], [assessment(first, relevance="irrelevant", reason="Its conditions do not match RQ1")], 3)["data"]
    assert out["selected_paper_ids"] == [second["paper_id"]]
    assert writing.get(created["project_id"], revision=2)["data"]["selected_paper_ids"] == [first["paper_id"]]


@pytest.mark.parametrize("change", ["fingerprint", "id", "basis", "reason", "score", "question", "duplicate"])
def test_assessment_invalid_shape_and_unknown_identifiers(writing, change):
    p = writing.library.put(record())
    created = writing.create(profile(), paper_ids=[p["paper_id"]])["data"]
    item = assessment(created["candidates"][0])
    if change == "fingerprint": item["metadata_sha256"] = "f" * 64
    elif change == "id": item["paper_id"] = paper_id_for("crossref", "unknown")
    elif change == "basis": item["basis"] = "abstract"
    elif change == "reason": item["reason"] = "   "
    elif change == "score": item["score"] = 0.99
    elif change == "question": item["question_ids"] = ["RQ99"]
    batch = [item, item] if change == "duplicate" else [item]
    with pytest.raises(JournalError):
        writing.assess(created["project_id"], batch, 1)
    assert writing.get(created["project_id"])["data"]["revision"] == 1


def test_fulltext_judgement_accepts_real_unselected_candidate_evidence(writing, tmp_path):
    created = local_project(writing, tmp_path, 2)
    row = created["evidence_matrix"][0]
    candidate = next(c for c in created["candidates"] if c["paper_id"] == row["paper_id"])
    selected = writing.assess(created["project_id"], [assessment(candidate, basis="fulltext", evidence_ids=[row["citation_id"]])],
                              1, selected_paper_ids=[])["data"]
    assert selected["selected_paper_ids"] == []
    assert selected["assessments"][0]["basis"] == "fulltext"
    with pytest.raises(JournalError) as exc:
        writing.save_draft(created["project_id"], draft_for(selected, row["citation_id"]), 2)
    assert exc.value.code == "UNSELECTED_EVIDENCE"


def test_fulltext_cross_paper_or_abstract_citation_rejected(writing, tmp_path):
    created = local_project(writing, tmp_path, 2)
    first, second = created["candidates"]
    row = next(r for r in created["evidence_matrix"] if r["paper_id"] == second["paper_id"])
    for evidence in [row["citation_id"], "cite-" + "f" * 24]:
        with pytest.raises(JournalError):
            writing.assess(created["project_id"], [assessment(first, basis="fulltext", evidence_ids=[evidence])], 1)
    with pytest.raises(JournalError):
        writing.assess(created["project_id"], [assessment(first, basis="fulltext", evidence_ids=[])], 1)


def test_get_all_candidates_prepare_only_selected_and_paginated(writing, tmp_path):
    created = local_project(writing, tmp_path, 2)
    first = created["candidates"][0]
    writing.assess(created["project_id"], [assessment(first)], 1)
    all_evidence = writing.get(created["project_id"], evidence_limit=1)
    assert all_evidence["data"]["evidence_total"] == 2
    assert all_evidence["data"]["next_evidence_offset"] == 1
    assert all_evidence["coverage"]["evidence_truncated"]
    last = writing.get(created["project_id"], evidence_offset=1, evidence_limit=1)
    assert last["data"]["next_evidence_offset"] is None
    assert last["data"]["evidence_matrix"][0]["citation_id"] != all_evidence["data"]["evidence_matrix"][0]["citation_id"]
    prepared = writing.prepare_writing(created["project_id"], evidence_limit=1)
    assert prepared["data"]["evidence_total"] == 1
    assert prepared["data"]["next_evidence_offset"] is None
    assert prepared["data"]["evidence_scope"] == "selected_papers"
    assert prepared["data"]["evidence_matrix"][0]["paper_id"] == first["paper_id"]
    assert "draft_schema" in prepared["data"]


def test_summary_and_abstract_are_never_evidence(writing):
    paper = writing.library.put(record(abstract="Abstract mentions excellent results"))
    created = writing.create(profile(), paper_ids=[paper["paper_id"]])["data"]
    selected = writing.assess(created["project_id"], [assessment(created["candidates"][0], basis="abstract")], 1)["data"]
    assert selected["evidence_matrix"] == []
    assert selected["stage"] == "needs_reading"


@pytest.mark.parametrize("tamper", ["quote", "page", "fragment", "sha", "paper", "bool_page"])
def test_saved_cards_rechecked_not_trusted_verbatim(writing, tmp_path, tamper):
    created = local_project(writing, tmp_path)
    pid = created["candidates"][0]["paper_id"]
    with writing.library.connection(True) as conn:
        row = conn.execute("SELECT * FROM readings WHERE paper_id=?", (pid,)).fetchone()
        card = json.loads(row["card"])
        evidence = card["claims"][0]["evidence"][0]
        if tamper == "quote": evidence["quote"] = "invented findings"
        elif tamper == "page": evidence["page_number"] = 99
        elif tamper == "fragment": evidence["fragment_id"] = "frag-" + "a" * 24
        elif tamper == "sha": card["file_sha256"] = "f" * 64
        elif tamper == "paper": card["paper_id"] = paper_id_for("crossref", "other")
        elif tamper == "bool_page": evidence["page_number"] = True
        conn.execute("UPDATE readings SET card=? WHERE reading_id=?", (json.dumps(card), row["reading_id"]))
    out = writing.get(created["project_id"])["data"]
    assert out["evidence_matrix"] == []
    assert any(g["code"] in ("EVIDENCE_MISMATCH", "STALE_EVIDENCE") for g in out["gaps"])


def test_fragment_locator_hash_is_recomputed(writing, tmp_path):
    created = local_project(writing, tmp_path)
    fragment = created["evidence_matrix"][0]["evidence"][0]["fragment_id"]
    with writing.library.connection(True) as conn:
        conn.execute("UPDATE fragments SET start=start+1 WHERE fragment_id=?", (fragment,))
    assert writing.get(created["project_id"])["data"]["evidence_total"] == 0


def test_unverified_cannot_supply_draft_or_fulltext_basis(writing, tmp_path):
    created = local_project(writing, tmp_path, kinds=["unverified"])
    citation = created["evidence_matrix"][0]["citation_id"]
    with pytest.raises(JournalError) as exc:
        writing.assess(created["project_id"], [assessment(created["candidates"][0], basis="fulltext", evidence_ids=[citation])], 1)
    assert exc.value.code == "UNVERIFIED_EVIDENCE"
    selected = writing.assess(created["project_id"], [assessment(created["candidates"][0])], 1)["data"]
    assert selected["stage"] == "needs_reading"
    with pytest.raises(JournalError) as exc:
        writing.save_draft(created["project_id"], draft_for(selected), 2)
    assert exc.value.code == "UNVERIFIED_EVIDENCE"


def test_agent_inference_is_not_presented_as_author_claim(writing, tmp_path):
    selected = selected_project(writing, tmp_path, kinds=["agent_inference"])
    with pytest.raises(JournalError) as exc:
        writing.save_draft(selected["project_id"], draft_for(selected), 2)
    assert exc.value.code == "INFERENCE_ATTRIBUTION"
    saved = writing.save_draft(selected["project_id"], draft_for(selected, kind="literature_inference"), 2)["data"]
    assert saved["draft_valid"] and saved["export_ready"]


def test_draft_and_outline_refs_extra_metadata_and_empty_citations_rejected(writing, tmp_path):
    selected = selected_project(writing, tmp_path)
    for change in ("outline", "metadata", "quote", "missing", "bool_level", "duplicate_sections"):
        draft = draft_for(selected)
        paragraph = draft["sections"][0]["paragraphs"][0]
        if change == "outline": draft["outline"][0]["evidence_ids"] = ["cite-" + "f" * 24]
        elif change == "metadata": draft["references"] = [{"doi": "10.fake/created"}]
        elif change == "quote": paragraph["quote"] = "invented evidence"
        elif change == "missing": paragraph["citation_ids"] = []
        elif change == "bool_level": draft["sections"][0]["level"] = True
        elif change == "duplicate_sections": draft["sections"].append(copy.deepcopy(draft["sections"][0]))
        with pytest.raises(JournalError):
            writing.save_draft(selected["project_id"], draft, 2)
    assert writing.get(selected["project_id"])["data"]["revision"] == 2


def test_draft_section_merges_preserve_previous_body(writing, tmp_path):
    selected = selected_project(writing, tmp_path)
    initial = draft_for(selected)
    initial["sections"].append({"id": "methods", "title": "Proposed method", "paragraphs": [
        {"text": "A proposed study, not completed experiments.", "kind": "author_proposal", "citation_ids": [], "own_material_ids": []}]})
    saved = writing.save_draft(selected["project_id"], initial, 2)["data"]
    replacement = {"sections": [{"id": "methods", "title": "Revised proposal", "paragraphs": [
        {"text": "A revised planned method.", "kind": "author_proposal"}]}]}
    updated = writing.save_draft(selected["project_id"], replacement, 3)["data"]
    assert updated["draft"]["sections"][0] == saved["draft"]["sections"][0]
    assert updated["draft"]["sections"][1]["paragraphs"][0]["text"] == "A revised planned method."
    assert writing.get(selected["project_id"], revision=3)["data"]["draft"] == saved["draft"]
    assert initial["sections"][1]["paragraphs"][0]["text"].startswith("A proposed")


@pytest.mark.parametrize("section", ["results", "experimental_results", "conclusion", "experiments", "研究结果", "结论"])
def test_without_own_experiments_results_remain_placeholder(writing, tmp_path, section):
    selected = selected_project(writing, tmp_path)
    draft = draft_for(selected)
    draft["sections"].append({"id": "own_results", "title": section,
                              "paragraphs": [{"text": "We measured a 99 percent improvement.", "kind": "author_proposal"}]})
    with pytest.raises(JournalError) as exc:
        writing.save_draft(selected["project_id"], draft, 2)
    assert exc.value.code == "OWN_MATERIAL_REQUIRED"
    draft["sections"][-1]["paragraphs"] = [{"text": "Await actual user experiments.", "kind": "placeholder"}]
    assert writing.save_draft(selected["project_id"], draft, 2)["data"]["draft_valid"]


def test_user_material_requires_actual_bound_profile_ids(writing, tmp_path):
    selected = selected_project(writing, tmp_path, own_materials=[{"id": "UM1", "text": "User supplied observed throughput: 12."}])
    draft = draft_for(selected)
    draft["sections"].append({"id": "results", "title": "User-supplied observations", "paragraphs": [
        {"text": "User supplied observed throughput: 12.", "kind": "user_material", "own_material_ids": ["UM99"]}]})
    with pytest.raises(JournalError) as exc:
        writing.save_draft(selected["project_id"], draft, 2)
    assert exc.value.code == "UNKNOWN_OWN_MATERIAL_ID"
    draft["sections"][-1]["paragraphs"][0]["own_material_ids"] = ["UM1"]
    out = writing.save_draft(selected["project_id"], draft, 2)
    assert out["data"]["export_ready"]
    assert not out["coverage"]["own_materials_independently_verified"]


def test_cas_stale_revision_never_calls_provider_or_overwrites(writing, tmp_path):
    selected = selected_project(writing, tmp_path)
    with pytest.raises(JournalError) as exc:
        writing.search(selected["project_id"], 1, queries=[query()])
    assert exc.value.code == "REVISION_CONFLICT"
    assert not writing.literature.providers.calls
    writing.save_draft(selected["project_id"], draft_for(selected), 2)
    with pytest.raises(JournalError) as exc:
        writing.save_draft(selected["project_id"], {"title": "Overwrite stale"}, 2)
    assert exc.value.code == "REVISION_CONFLICT"
    assert writing.get(selected["project_id"])["data"]["draft"]["title"] != "Overwrite stale"


@pytest.mark.parametrize("method,kwargs", [
    ("list", {"limit": True}), ("list", {"offset": False}), ("list", {"limit": 101}),
    ("get", {"revision": True}), ("get", {"evidence_limit": False}), ("get", {"evidence_offset": -1}),
    ("prepare_writing", {"evidence_limit": 101}),
])
def test_bool_and_bad_integer_arguments_are_journal_errors(writing, method, kwargs):
    created = writing.create(profile())["data"]
    with pytest.raises(JournalError):
        getattr(writing, method)(**kwargs) if method == "list" else getattr(writing, method)(created["project_id"], **kwargs)


def test_two_mib_budget_and_error_no_raw_input(writing):
    with pytest.raises(JournalError) as exc:
        writing.prepare("x" * (MAX_INPUT_BYTES + 1))
    assert exc.value.code == "INPUT_TOO_LARGE"
    with pytest.raises(JournalError) as exc:
        writing.create(profile(), source_text="SECRET" * MAX_INPUT_BYTES)
    assert exc.value.code == "INPUT_TOO_LARGE" and "SECRET" not in str(exc.value)
    assert not writing.library.directory.exists()


def test_render_and_export_delayed_validated_payload_only(writing, tmp_path, monkeypatch):
    selected = selected_project(writing, tmp_path)
    out = writing.save_draft(selected["project_id"], draft_for(selected), 2)["data"]
    payloads = fake_export(monkeypatch)
    assert writing.render_docx(selected["project_id"]) == b"verified docx fixture"
    result = writing.export_docx(selected["project_id"], str(tmp_path / "draft.docx"))
    assert result["data"]["revision"] == 3
    payload = payloads[0]
    assert set(payload) == {"project_id", "revision", "profile", "draft", "citations", "references", "gaps", "coverage"}
    assert set(payload["citations"]) == {out["evidence_matrix"][0]["citation_id"]}
    assert payload["references"][0]["paper_id"] == selected["selected_paper_ids"][0]
    assert payload["coverage"]["evidence_check_only"] and not payload["coverage"]["semantic_correctness_verified"]
    assert not writing.literature.providers.calls


def test_export_without_body_evidence_and_mere_outline_cites_fails(writing, tmp_path, monkeypatch):
    selected = selected_project(writing, tmp_path)
    payloads = fake_export(monkeypatch)
    with pytest.raises(JournalError) as exc:
        writing.render_docx(selected["project_id"])
    assert exc.value.code == "DRAFT_REQUIRED"
    draft = draft_for(selected)
    draft["sections"][0]["paragraphs"] = [{"text": "Await writing", "kind": "placeholder"}]
    writing.save_draft(selected["project_id"], draft, 2)
    with pytest.raises(JournalError) as exc:
        writing.render_docx(selected["project_id"])
    assert exc.value.code == "CITED_EVIDENCE_REQUIRED" and not payloads


def test_hash_change_rechecked_at_get_save_render_and_export(writing, tmp_path, monkeypatch):
    selected = selected_project(writing, tmp_path)
    draft = draft_for(selected)
    writing.save_draft(selected["project_id"], draft, 2)
    payloads = fake_export(monkeypatch)
    sha = selected["evidence_matrix"][0]["file_sha256"]
    writing.library.pdf_path(sha).write_bytes(b"changed PDF bytes")
    current = writing.get(selected["project_id"])["data"]
    assert not current["export_ready"] and not current["draft_valid"] and current["evidence_matrix"] == []
    for action in (lambda: writing.save_draft(selected["project_id"], {"title": "re-save"}, 3),
                   lambda: writing.render_docx(selected["project_id"]),
                   lambda: writing.export_docx(selected["project_id"], str(tmp_path / "draft.docx"))):
        with pytest.raises(JournalError): action()
    assert not payloads


def test_metadata_fingerprint_ignores_retrieval_but_not_real_changes(writing, tmp_path, monkeypatch):
    selected = selected_project(writing, tmp_path)
    writing.save_draft(selected["project_id"], draft_for(selected), 2)
    candidate = selected["candidates"][0]
    paper = copy.deepcopy(candidate["paper"])
    paper.update(retrieved_at="2099-01-01T00:00:00Z", query="unrelated later search")
    writing.library.put(paper)
    payloads = fake_export(monkeypatch)
    assert writing.render_docx(selected["project_id"])
    paper["title"] = "Actually different metadata"
    writing.library.put(paper)
    current = writing.get(selected["project_id"])["data"]
    assert current["candidates"][0]["metadata_sha256"] != candidate["metadata_sha256"]
    assert not current["draft_valid"]
    with pytest.raises(JournalError) as exc:
        writing.render_docx(selected["project_id"])
    assert exc.value.code == "METADATA_MISMATCH"
    assert len(payloads) == 1


def test_exporter_failure_has_only_journal_error_not_secrets(writing, tmp_path, monkeypatch):
    selected = selected_project(writing, tmp_path)
    writing.save_draft(selected["project_id"], draft_for(selected), 2)
    fake_export(monkeypatch, fail=RuntimeError("rawtrace token=SECRET"))
    with pytest.raises(JournalError) as exc:
        writing.render_docx(selected["project_id"])
    assert exc.value.code == "WRITING_EXPORT_ERROR"
    assert "SECRET" not in str(exc.value) and "rawtrace" not in str(exc.value)


@pytest.mark.parametrize("title,kind,text", [
    ("相关文献的实验结果", "literature_summary", "已有论文的作者描述了可复现方法。"),
    ("结论", "literature_summary", "已有文献的作者描述的方法提供了研究背景。"),
    ("文献支持的结论", "literature_inference", "据此可以推断，既有方法可能支持未来研究。"),
    ("实验计划", "author_proposal", "我们计划在未来开展实验，目前尚未完成。"),
    ("experiment_plan", "author_proposal", "We plan future experiments; no experiments have yet been completed."),
])
def test_review_results_conclusions_and_unfinished_plans_are_allowed(writing, tmp_path, title, kind, text):
    selected = selected_project(writing, tmp_path)
    draft = draft_for(selected)
    paragraph = {"text": text, "kind": kind}
    if kind.startswith("literature_"):
        paragraph["citation_ids"] = [selected["evidence_matrix"][0]["citation_id"]]
    draft["sections"].append({"id": "review_conclusion", "title": title, "paragraphs": [paragraph]})
    assert writing.save_draft(selected["project_id"], draft, 2)["data"]["export_ready"]


@pytest.mark.parametrize("text", ["我们测得性能提升了百分之九十九。", "本研究实验结果表明性能已提高。", "We obtained a 99 percent improvement."])
def test_unsupported_own_completed_results_rejected_even_under_methods(writing, tmp_path, text):
    selected = selected_project(writing, tmp_path)
    draft = draft_for(selected)
    draft["sections"].append({"id": "methods", "title": "Methods", "paragraphs": [{"text": text, "kind": "author_proposal"}]})
    with pytest.raises(JournalError) as exc:
        writing.save_draft(selected["project_id"], draft, 2)
    assert exc.value.code == "OWN_MATERIAL_REQUIRED"


@pytest.mark.parametrize("field,value", [("title", "Different outline title"), ("level", 2)])
def test_outline_and_section_same_id_mismatch_rejected_early(writing, tmp_path, field, value):
    selected = selected_project(writing, tmp_path)
    draft = draft_for(selected)
    draft["outline"][0][field] = value
    with pytest.raises(JournalError) as exc:
        writing.save_draft(selected["project_id"], draft, 2)
    assert exc.value.code == "CITATION_MISMATCH"
    assert writing.get(selected["project_id"])["data"]["revision"] == 2


def test_empty_placeholder_is_legal_and_real_docx_supplies_notice(writing, tmp_path):
    import io
    from docx import Document
    selected = selected_project(writing, tmp_path)
    draft = draft_for(selected)
    draft["sections"].append({"id": "results", "title": "Actual user results", "paragraphs": [{"text": "", "kind": "placeholder"}]})
    saved = writing.save_draft(selected["project_id"], draft, 2)["data"]
    assert saved["draft"]["sections"][-1]["paragraphs"][0]["text"] == ""
    raw = writing.render_docx(selected["project_id"])
    text = "\n".join(p.text for p in Document(io.BytesIO(raw)).paragraphs)
    assert "待补" in text or "placeholder" in text.lower() or "pending" in text.lower()
    assert writing.export_docx(selected["project_id"], str(tmp_path / "actual-draft.docx"))["data"]["revision"] == 3
    assert (tmp_path / "actual-draft.docx").read_bytes().startswith(b"PK")
    with pytest.raises(JournalError):
        writing.export_docx(selected["project_id"], str(tmp_path / "actual-draft.docx"))


@pytest.mark.parametrize("kind", ["literature_summary", "literature_inference", "author_proposal", "user_material"])
def test_empty_nonplaceholder_is_not_a_body_paragraph(writing, tmp_path, kind):
    selected = selected_project(writing, tmp_path)
    draft = draft_for(selected, kind=kind)
    draft["sections"][0]["paragraphs"][0]["text"] = ""
    with pytest.raises(JournalError) as exc:
        writing.save_draft(selected["project_id"], draft, 2)
    assert exc.value.code == "INVALID_DRAFT"


@pytest.mark.parametrize("token", ["[1]", "［1］", "[@FAKEKEY]", "@FabricatedZoteroKey", "ZOTERO_ITEM"])
def test_manual_number_or_zotero_marker_rejected_at_save(writing, tmp_path, token):
    selected = selected_project(writing, tmp_path)
    draft = draft_for(selected)
    draft["sections"][0]["paragraphs"][0]["text"] += " " + token
    with pytest.raises(JournalError) as exc:
        writing.save_draft(selected["project_id"], draft, 2)
    assert exc.value.code == "INVALID_DRAFT"


def test_fragment_and_quote_joint_tampering_does_not_pass_pdf_recheck(writing, tmp_path):
    created = local_project(writing, tmp_path)
    evidence = created["evidence_matrix"][0]
    fid = evidence["evidence"][0]["fragment_id"]
    with writing.library.connection(True) as conn:
        fragment = conn.execute("SELECT text FROM fragments WHERE fragment_id=?", (fid,)).fetchone()[0]
        fake_text = "X" * len(fragment)
        conn.execute("UPDATE fragments SET text=? WHERE fragment_id=?", (fake_text, fid))
        row = conn.execute("SELECT card FROM readings WHERE reading_id=?", (evidence["reading_id"],)).fetchone()
        card = json.loads(row[0])
        card["claims"][0]["evidence"][0]["quote"] = fake_text[:20]
        import hashlib
        from paperflow.engine.literature.writing_models import budget_json
        payload = budget_json(card)
        new_id = "reading-" + hashlib.sha256((evidence["paper_id"] + evidence["file_sha256"] + "calling_agent" + payload).encode("utf-8")).hexdigest()[:24]
        conn.execute("UPDATE readings SET card=?,reading_id=? WHERE reading_id=?", (payload, new_id, evidence["reading_id"]))
    project = writing.get(created["project_id"])["data"]
    assert not project["evidence_matrix"]
    assert any(g["code"] == "EVIDENCE_MISMATCH" for g in project["gaps"])


def test_pdf_recheck_caches_reader_and_page_per_request(writing, tmp_path, monkeypatch):
    import pypdf
    selected = selected_project(writing, tmp_path)
    pid = selected["selected_paper_ids"][0]
    card = card_for(writing.literature.read(pid))
    card["claims"] = [copy.deepcopy(card["claims"][0]) for _ in range(5)]
    for number, claim in enumerate(card["claims"]):
        claim["text"] += " " + str(number)
    writing.literature.save_reading(pid, card)
    real_reader, calls = pypdf.PdfReader, []
    def counting_reader(*args, **kwargs):
        calls.append(1)
        return real_reader(*args, **kwargs)
    monkeypatch.setattr(pypdf, "PdfReader", counting_reader)
    current = writing.get(selected["project_id"])["data"]
    assert current["evidence_total"] == 6 and len(calls) == 1
    writing.save_draft(selected["project_id"], draft_for(selected), 2)
    assert len(calls) == 2  # Save + response reuse the same request-local verifier.


def test_filesystem_failure_does_not_leak_paths_or_rawtrace(writing, monkeypatch):
    def fail(*args, **kwargs):
        raise PermissionError("private path APIKEY=SECRET rawtrace")
    monkeypatch.setattr(writing.store, "create", fail)
    with pytest.raises(JournalError) as exc:
        writing.create(profile())
    assert exc.value.code == "WRITING_STORE_ERROR"
    assert "SECRET" not in str(exc.value) and "private" not in str(exc.value)


def test_real_docx_reference_and_evidence_appendix_are_backend_derived(writing, tmp_path):
    import io
    from docx import Document
    selected = selected_project(writing, tmp_path)
    writing.save_draft(selected["project_id"], draft_for(selected), 2)
    raw = writing.render_docx(selected["project_id"])
    document = Document(io.BytesIO(raw))
    text = "\n".join(p.text for p in document.paragraphs)
    text += "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "[1]" in text
    assert selected["candidates"][0]["paper"]["title"] in text
    assert selected["evidence_matrix"][0]["reading_id"] in text
    assert selected["evidence_matrix"][0]["file_sha256"] in text
    assert not writing.literature.providers.calls


def test_changed_query_plan_preserves_history_without_unknown_match_ids(writing):
    created = writing.create(profile(), queries=[query(i) for i in range(1, 7)])["data"]
    first = writing.search(created["project_id"], 1)["data"]
    assert first["candidates"][0]["matched_query_ids"] == ["Q" + str(i) for i in range(1, 7)]
    second = writing.search(created["project_id"], 2, queries=[query(i) for i in range(11, 17)])["data"]
    assert second["candidates"][0]["matched_query_ids"] == ["Q" + str(i) for i in range(11, 17)]
    history = writing.get(created["project_id"], revision=2)["data"]
    assert history["candidates"][0]["matched_query_ids"] == first["candidates"][0]["matched_query_ids"]


def test_candidate_cap_does_not_fake_new_results_or_delete_old_body(writing):
    ids = [writing.library.put(record(identifier="seed-" + str(i), doi="10.1234/seed-" + str(i)))["paper_id"] for i in range(60)]
    created = writing.create(profile(), paper_ids=ids, queries=[query()])["data"]
    result = writing.search(created["project_id"], 1)
    assert result["status"] == "partial"
    assert len(result["data"]["candidates"]) == 60
    assert {c["paper_id"] for c in result["data"]["candidates"]} == set(ids)
    assert result["warnings"]
    assert writing.library.list()["total"] == 60


def test_stale_older_judgement_does_not_block_valid_partial_batch(writing, tmp_path):
    selected = selected_project(writing, tmp_path, count=2)
    first, second = selected["candidates"]
    changed = copy.deepcopy(first["paper"])
    changed["title"] = "Different currently stored title"
    writing.library.put(changed)
    saved = writing.assess(selected["project_id"], [assessment(second, relevance="background")], 2)["data"]
    assert len(saved["assessments"]) == 2
    assert saved["selected_paper_ids"] == [second["paper_id"]]
    assert any(g["code"] == "METADATA_MISMATCH" for g in saved["gaps"])


def test_draft_language_is_inherited_and_explicit_mismatch_rejected(writing, tmp_path):
    local = local_project(writing, tmp_path)
    created = writing.create(profile(language="en"), paper_ids=[local["candidates"][0]["paper_id"]])["data"]
    selected = writing.assess(created["project_id"], [assessment(created["candidates"][0])], 1)["data"]
    draft = draft_for(selected)
    draft["language"] = "zh"
    with pytest.raises(JournalError) as exc:
        writing.save_draft(selected["project_id"], draft, 2)
    assert exc.value.code == "LANGUAGE_MISMATCH"
    draft.pop("language")
    saved = writing.save_draft(selected["project_id"], draft, 2)["data"]
    assert saved["draft"]["language"] == "en"
    assert writing.render_docx(selected["project_id"]).startswith(b"PK")


def test_spoofed_source_identity_is_reported_not_saved(writing):
    malformed = record()
    malformed["paper_id"] = "paper-" + "f" * 24
    writing.literature.providers.records = {"crossref": [malformed], "arxiv": []}
    created = writing.create(profile(), queries=[query()])["data"]
    result = writing.search(created["project_id"], 1)
    assert result["status"] == "partial"
    assert result["data"]["candidates"] == []
    assert result["data"]["search_statuses"][0]["error_code"] == "INVALID_SOURCE_DATA"
    assert writing.library.list()["total"] == 0


def test_reading_identity_hash_prevents_silent_claim_replacement(writing, tmp_path):
    created = local_project(writing, tmp_path)
    original = created["evidence_matrix"][0]
    with writing.library.connection(True) as conn:
        raw = conn.execute("SELECT card FROM readings WHERE reading_id=?", (original["reading_id"],)).fetchone()[0]
        card = json.loads(raw)
        card["claims"][0]["text"] = "A different substituted author assertion"
        conn.execute("UPDATE readings SET card=? WHERE reading_id=?", (json.dumps(card), original["reading_id"]))
    data = writing.get(created["project_id"])["data"]
    assert not data["evidence_matrix"]
    assert any(g["code"] == "EVIDENCE_MISMATCH" for g in data["gaps"])


def test_core_sources_parse_with_python39_grammar():
    import ast
    root = Path(__file__).resolve().parents[1]
    for filename in ("writing_models.py", "writing_store.py", "writing_service.py"):
        file = root / "paperflow" / "engine" / "literature" / filename
        ast.parse(file.read_text(encoding="utf-8"), feature_version=(3, 9))
