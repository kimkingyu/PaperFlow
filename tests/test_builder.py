"""Unit tests for PaperFlow docx_builder and anti_ai_cleaner.

Run:  .venv\\Scripts\\python.exe -m unittest tests/test_builder.py -v
These tests never touch a running Word instance.
"""

import os
import tempfile
import unittest

from docx import Document
from docx.shared import Pt

from paperflow.engine.anti_ai_cleaner import anti_ai_engine
from paperflow.engine.docx_builder import AcademicDocxBuilder


class TestAntiAI(unittest.TestCase):
    def test_detects_cliches(self) -> None:
        text = "不可否认的是，在当今飞速发展的背景下，本研究旨在深入探讨深度学习算法。"
        result = anti_ai_engine.analyze(text)
        self.assertGreater(result["total_findings"], 0)


class TestDocxBuilder(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.path = os.path.join(self.tmp.name, "paper.docx")

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def _build(self) -> Document:
        b = AcademicDocxBuilder(is_chinese=True)
        b.add_title("基于动态量化的轻量级边缘推理研究")
        b.add_abstract("本文提出了一种自适应量化算法。", ["边缘计算", "模型量化"])
        b.add_heading("1 引言", level=1)
        b.add_heading("1.1 研究背景", level=2)
        b.add_paragraph_with_citations("大规模模型面临显存压力 [@vaswani2017, @he2016]，已有工作 [@vaswani2017]。")
        b.add_bibliography_section()
        b.save(self.path)
        return Document(self.path)

    def test_file_is_written(self) -> None:
        self._build()
        self.assertTrue(os.path.getsize(self.path) > 0)

    def test_headings_use_builtin_styles(self) -> None:
        doc = self._build()
        styles = {p.text: p.style.name for p in doc.paragraphs}
        self.assertEqual(styles["1 引言"], "Heading 1")
        self.assertEqual(styles["1.1 研究背景"], "Heading 2")

    def test_heading_sizes_match_preset(self) -> None:
        doc = self._build()
        self.assertEqual(doc.styles["Heading 1"].font.size, Pt(16))
        self.assertEqual(doc.styles["Heading 2"].font.size, Pt(14))
        self.assertEqual(doc.styles["Normal"].font.size, Pt(12))

    def test_citations_become_zotero_fields_with_stable_numbers(self) -> None:
        doc = self._build()
        xml = doc.element.body.xml
        self.assertIn("ADDIN ZOTERO_ITEM CSL_CITATION", xml)
        self.assertIn("ADDIN ZOTERO_BIBL", xml)
        # vaswani2017 first seen -> [1]; he2016 -> [2]; reuse keeps [1]
        self.assertIn("[1,2]", xml)
        self.assertEqual(xml.count(">[1]<"), 1)


    def test_three_line_table_structure(self) -> None:
        b = AcademicDocxBuilder(is_chinese=True)
        b.add_three_line_table(
            headers=["算法", "延迟(ms)", "准确率"],
            rows=[["FP32", "120.0", "85.2%"], ["INT8", "35.1", "84.9%"]],
            caption="表 1-1  实验对比",
        )
        b.save(self.path)
        doc = Document(self.path)
        self.assertEqual(len(doc.tables), 1)
        tbl = doc.tables[0]
        self.assertEqual(len(tbl.rows), 3)
        self.assertEqual(len(tbl.columns), 3)
        self.assertEqual(tbl.rows[0].cells[0].text, "算法")
        self.assertEqual(tbl.rows[1].cells[0].text, "FP32")
        xml = tbl._tbl.tblPr.xml
        self.assertIn('w:top w:val="single"', xml)
        self.assertIn('w:bottom w:val="single"', xml)
        # Verify caption
        self.assertTrue(any("表 1-1  实验对比" in p.text for p in doc.paragraphs))


if __name__ == "__main__":
    unittest.main()
