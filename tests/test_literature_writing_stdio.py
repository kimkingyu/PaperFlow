"""Actual MCP stdio, using generated PDF fixtures only; no research, network, models or live Word."""
from __future__ import annotations

import asyncio
import copy
import json
import os
import sys
import textwrap
from pathlib import Path

import pytest
from docx import Document
from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client


ROOT = Path(__file__).resolve().parent.parent
WRITING_TOOLS = {
    "prepare_paper_writing_research", "create_paper_writing_project", "search_related_papers",
    "assess_related_papers", "get_paper_writing_project", "list_paper_writing_projects",
    "prepare_paper_manuscript", "save_paper_manuscript", "export_paper_manuscript",
}
READ_ONLY = {"prepare_paper_writing_research", "get_paper_writing_project",
             "list_paper_writing_projects", "prepare_paper_manuscript"}
ENVELOPE_KEYS = {"status", "error_code", "message", "data", "sources", "coverage", "warnings", "suggested_options"}

# Guards run inside the real server subprocess, not only in pytest's process.
OFFLINE_SERVER = textwrap.dedent("""
    import socket
    import sys
    from importlib.abc import MetaPathFinder

    original_connect = socket.socket.connect
    original_getaddrinfo = socket.getaddrinfo

    def blocked(*args, **kwargs):
        raise AssertionError("offline stdio fixture must not contact providers, models or live Word")

    def local_socketpair_only(sock, address):
        if isinstance(address, tuple) and address[0] in ("127.0.0.1", "::1", "localhost"):
            return original_connect(sock, address)
        return blocked()

    def local_resolution_only(host, *args, **kwargs):
        if host in ("127.0.0.1", "::1", "localhost", None):
            return original_getaddrinfo(host, *args, **kwargs)
        return blocked()

    socket.create_connection = blocked
    socket.socket.connect = local_socketpair_only
    socket.getaddrinfo = local_resolution_only

    class NoModelsOrGui(MetaPathFinder):
        def find_spec(self, fullname, path=None, target=None):
            prefixes = ("paperflow.gui", "openai", "anthropic", "google.generativeai")
            if any(fullname == prefix or fullname.startswith(prefix + ".") for prefix in prefixes):
                raise AssertionError("stdio writing must not initialize a GUI or external model")
            return None

    sys.meta_path.insert(0, NoModelsOrGui())
    from paperflow.engine.word_live_bridge import live_bridge
    live_bridge.connect = blocked
    live_bridge._call = blocked
    from paperflow import __main__ as entrypoint
    sys.argv = ["paperflow", "run"]
    entrypoint.main()
""")


def _fixture_pdf(tmp_path, suffix):
    pypdf = pytest.importorskip("pypdf")
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = pypdf.PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject()
    sentence = f"Evidence fixture {suffix}: authors describe a bounded local protocol test. ".encode("ascii")
    stream.set_data(b"BT /F1 12 Tf 20 760 Td (" + sentence * 35 + b") Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.add_metadata({"/Title": f"Offline protocol fixture {suffix}"})
    target = tmp_path / f"offline-fixture-{suffix}.pdf"
    with target.open("wb") as output:
        writer.write(output)
    return target


def _document_text(document):
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_paper_writing_local_pdf_evidence_and_cited_docx_over_real_stdio(tmp_path):
    pdf_a = _fixture_pdf(tmp_path, "A")
    pdf_b = _fixture_pdf(tmp_path, "B")

    async def workflow():
        paper_home = tmp_path / "papers"
        research_home = tmp_path / "research"
        journal_home = tmp_path / "journals"
        params = StdioServerParameters(
            command=sys.executable, args=["-B", "-c", OFFLINE_SERVER], cwd=str(ROOT),
            env={**os.environ, "PYTHONPATH": str(ROOT), "PYTHONIOENCODING": "utf-8",
                 "PYTHONDONTWRITEBYTECODE": "1", "PAPERFLOW_PAPER_HOME": str(paper_home),
                 "PAPERFLOW_RESEARCH_HOME": str(research_home), "PAPERFLOW_JOURNAL_HOME": str(journal_home)},
        )
        async with stdio_client(params) as (read, write):
            async with ClientSession(read, write) as session:
                await session.initialize()
                tools = {item.name: item for item in (await session.list_tools()).tools}
                assert WRITING_TOOLS <= tools.keys()
                assert {"import_local_paper", "read_academic_paper", "save_paper_reading",
                        "create_research_project", "select_target_word_doc", "recommend_journals"} <= tools.keys()
                for name in WRITING_TOOLS:
                    annotations = tools[name].annotations.model_dump(by_alias=True)
                    assert annotations["readOnlyHint"] is (name in READ_ONLY)
                    assert annotations["openWorldHint"] is (name == "search_related_papers")
                    assert annotations["destructiveHint"] is (name == "export_paper_manuscript")
                for name in ("search_related_papers", "assess_related_papers", "save_paper_manuscript"):
                    schema = tools[name].model_dump(by_alias=True)["inputSchema"]
                    assert "expected_revision" in schema["required"]

                async def call(name, arguments):
                    result = await session.call_tool(name, arguments)
                    assert not result.model_dump(by_alias=True).get("isError", False), result
                    payload = json.loads("\n".join(item.text for item in result.content if hasattr(item, "text")))
                    if name in WRITING_TOOLS:
                        assert set(payload) == ENVELOPE_KEYS
                    else:
                        # Existing literature success responses intentionally omit
                        # optional error_code/message; keep their old semantics.
                        assert {"status", "data", "sources", "coverage", "warnings", "suggested_options"} <= payload.keys()
                    return payload

                async def success(name, arguments):
                    payload = await call(name, arguments)
                    assert payload["status"] in ("success", "partial"), payload
                    return payload

                empty = await success("list_paper_writing_projects", {})
                assert empty["data"]["total"] == 0 and empty["data"]["projects"] == []
                prepared = await success("prepare_paper_writing_research", {
                    "text": "离线stdio协议测试fixture：核对文献出处边界，不产生实际研究或实验结果。"})
                for schema in ("profile_schema", "query_schema", "assessment_schema", "draft_schema"):
                    assert "properties" in prepared["data"][schema]
                assert prepared["coverage"]["backend_calls_llm"] is False
                assert not paper_home.exists() and not research_home.exists() and not journal_home.exists()
                malformed = await session.call_tool("get_paper_writing_project", {"project_id": "fixture", "revision": True})
                assert malformed.model_dump(by_alias=True).get("isError", False)
                assert not paper_home.exists()

                imported_a = await success("import_local_paper", {"file_path": str(pdf_a), "title": "Offline protocol fixture A"})
                imported_b = await success("import_local_paper", {"file_path": str(pdf_b), "title": "Offline protocol fixture B"})
                paper_a, paper_b = imported_a["data"]["paper_id"], imported_b["data"]["paper_id"]
                assert paper_a != paper_b
                created = await success("create_paper_writing_project", {
                    "profile": {"title": "MCP文献支持草稿fixture", "research_question": "离线接口是否保留逐页引用来源？",
                                "language": "zh", "sub_questions": [{"id": "RQ1", "question": "引用能否定位当前片段？"}],
                                "own_materials": []},
                    "queries": [{"id": "Q1", "query": "offline bounded protocol fixture", "purpose": "测试检索计划存储而不执行网络",
                                 "question_ids": ["RQ1"], "sources": ["arxiv", "crossref"]}],
                    "paper_ids": [paper_a, paper_b], "source_text": prepared["data"]["text"],
                    "input_id": prepared["data"]["input_id"],
                })
                project_id = created["data"]["project_id"]
                assert created["data"]["revision"] == 1
                candidates = {candidate["paper_id"]: candidate for candidate in created["data"]["candidates"]}
                assert set(candidates) == {paper_a, paper_b}
                assert candidates[paper_a]["paper"]["title"] == "Offline protocol fixture A"
                assessments = [
                    {"paper_id": paper_a, "metadata_sha256": candidates[paper_a]["metadata_sha256"],
                     "relevance": "core", "reason": "主fixture用于核对正文引用协议，不是领域相关性认证", "question_ids": ["RQ1"],
                     "basis": "metadata", "limitations": ["只看fixture元数据，尚未读正文"], "evidence_ids": []},
                    {"paper_id": paper_b, "metadata_sha256": candidates[paper_b]["metadata_sha256"],
                     "relevance": "marginal", "reason": "备用fixture只检查未选候选的矩阵覆盖", "question_ids": ["RQ1"],
                     "basis": "metadata", "limitations": ["不纳入本次草稿"], "evidence_ids": []},
                ]
                assessed = await success("assess_related_papers", {"project_id": project_id, "assessments": assessments,
                                                                   "expected_revision": 1, "selected_paper_ids": [paper_a]})
                assert assessed["data"]["revision"] == 2 and assessed["data"]["selected_paper_ids"] == [paper_a]
                assert assessed["data"]["stage"] == "needs_reading"
                before_reading = await success("prepare_paper_manuscript", {"project_id": project_id})
                assert before_reading["data"]["evidence_matrix"] == []

                cards = {}
                for paper_id in (paper_a, paper_b):
                    first = await success("read_academic_paper", {"paper_id": paper_id, "max_chars": 1000})
                    first_fragment = first["data"]["fragments"][0]
                    sha = first["data"]["file_sha256"]
                    assert first["data"]["next_cursor"] is not None
                    assert first["coverage"]["understanding_complete"] is False
                    cursor = first["data"]["next_cursor"]
                    final = first
                    while cursor is not None:
                        final = await success("read_academic_paper", {
                            "paper_id": paper_id, "page_number": cursor["page_number"], "offset": cursor["offset"],
                            "page_count": 3, "max_chars": 1000})
                        assert final["data"]["file_sha256"] == sha
                        cursor = final["data"]["next_cursor"]
                    assert final["coverage"]["text_complete"] is True
                    assert final["coverage"]["understanding_complete"] is False
                    card = {"paper_id": paper_id, "file_sha256": sha, "summary": "仅解释离线测试fixture，不是实际研究结果。",
                            "claims": [{"section": "method", "kind": "author_claim",
                                        "text": "fixture原文描述了有界本地协议测试，仅用于出处检查。",
                                        "evidence": [{"fragment_id": first_fragment["fragment_id"],
                                                      "page_number": first_fragment["page_number"],
                                                      "quote": first_fragment["text"][:150]}]}]}
                    saved_card = await success("save_paper_reading", {"paper_id": paper_id, "reading": card})
                    assert saved_card["data"]["evidence_check_only"] is True
                    assert saved_card["data"]["semantic_correctness_verified"] is False
                    cards[paper_id] = {"card": card, "reading_id": saved_card["data"]["reading_id"]}

                all_evidence = await success("get_paper_writing_project", {
                    "project_id": project_id, "evidence_offset": 0, "evidence_limit": 1})
                assert all_evidence["data"]["evidence_total"] == 2, all_evidence["data"]["gaps"]
                assert all_evidence["data"]["evidence_scope"] == "all_candidates"
                assert all_evidence["data"]["next_evidence_offset"] == 1
                next_evidence = await success("get_paper_writing_project", {
                    "project_id": project_id, "evidence_offset": all_evidence["data"]["next_evidence_offset"], "evidence_limit": 1})
                rows = all_evidence["data"]["evidence_matrix"] + next_evidence["data"]["evidence_matrix"]
                assert {row["paper_id"] for row in rows} == {paper_a, paper_b}
                assert next_evidence["data"]["next_evidence_offset"] is None
                writing = await success("prepare_paper_manuscript", {"project_id": project_id, "evidence_limit": 1})
                assert writing["data"]["evidence_total"] == 1 and writing["data"]["evidence_scope"] == "selected_papers"
                assert writing["data"]["next_evidence_offset"] is None
                assert "draft_schema" in writing["data"] and "agent_contract" in writing["data"]
                row = writing["data"]["evidence_matrix"][0]
                assert row["paper_id"] == paper_a and row["reading_id"] == cards[paper_a]["reading_id"]
                assert row["file_sha256"] == cards[paper_a]["card"]["file_sha256"]
                original_evidence = cards[paper_a]["card"]["claims"][0]["evidence"]
                assert len(row["evidence"]) == len(original_evidence)
                for actual, original in zip(row["evidence"], original_evidence):
                    assert {key: actual[key] for key in original} == original
                    assert actual.get("match_method", "exact") == "exact"
                citation_id = row["citation_id"]

                fulltext_assessment = {**assessments[0], "basis": "fulltext", "evidence_ids": [citation_id],
                                       "reason": "fixture已读片段支持本次出处协议测试，不是科学结论", "limitations": ["未认证语义正确性"]}
                upgraded = await success("assess_related_papers", {"project_id": project_id, "assessments": [fulltext_assessment],
                                                                   "expected_revision": 2})
                assert upgraded["data"]["revision"] == 3
                assert upgraded["data"]["selected_paper_ids"] == [paper_a]
                assert len(upgraded["data"]["assessments"]) == 2

                draft = {"title": "有出处的离线fixture草稿", "language": "zh",
                         "outline": [{"id": "intro", "title": "引言", "level": 1, "purpose": "引用真实fixture片段", "evidence_ids": [citation_id]},
                                     {"id": "results", "title": "待补结果", "level": 1, "purpose": "没有用户实验数据", "evidence_ids": []}],
                         "sections": [{"id": "intro", "title": "引言", "level": 1,
                                       "paragraphs": [{"text": "离线夹具的原文描述了有界本地协议测试；这不是实际研究结果。",
                                                       "kind": "literature_summary", "citation_ids": [citation_id], "own_material_ids": []}]},
                                      {"id": "results", "title": "待补结果", "level": 1,
                                       "paragraphs": [{"text": "未提供用户真实实验数据；结果与结论待补。", "kind": "placeholder",
                                                       "citation_ids": [], "own_material_ids": []}]}]}
                bad_draft = copy.deepcopy(draft)
                bad_draft["sections"][0]["paragraphs"][0]["citation_ids"] = []
                rejected = await call("save_paper_manuscript", {"project_id": project_id, "draft": bad_draft,
                                                               "expected_revision": 3, "change_note": "测试拒绝无出处正文"})
                assert rejected["status"] == "error"
                stored = await success("save_paper_manuscript", {"project_id": project_id, "draft": draft,
                                                                 "expected_revision": 3, "change_note": "fixture编号引用正文"})
                assert stored["data"]["revision"] == 4
                assert stored["data"]["draft"]["sections"][0]["paragraphs"][0]["citation_ids"] == [citation_id]
                assert stored["coverage"]["evidence_check_only"] is True
                assert stored["coverage"]["semantic_correctness_verified"] is False
                conflict = await call("save_paper_manuscript", {"project_id": project_id, "draft": draft,
                                                               "expected_revision": 3, "change_note": "旧版本拒绝覆盖"})
                assert conflict["error_code"] == "REVISION_CONFLICT"
                historical = await success("get_paper_writing_project", {"project_id": project_id, "revision": 1})
                assert historical["data"]["revision"] == 1 and historical["data"]["draft"] is None

                destination = tmp_path / "cited-fixture.docx"
                exported = await success("export_paper_manuscript", {"project_id": project_id, "output_path": str(destination)})
                assert exported["data"]["revision"] == 4
                document = Document(destination)
                text = _document_text(document)
                assert "[1]" in text and "Offline protocol fixture A" in text
                assert "Offline protocol fixture B" not in text
                assert "未提供用户真实实验数据" in text
                assert citation_id in text and row["file_sha256"] in text
                assert row["evidence"][0]["fragment_id"] in text and row["reading_id"] in text
                # DOCX carries source locators, not automatic copies of PDF quotes.
                assert row["evidence"][0]["quote"] not in text
                assert "ADDIN ZOTERO" not in document.element.xml
                original_bytes = destination.read_bytes()
                kept = await call("export_paper_manuscript", {"project_id": project_id, "output_path": str(destination)})
                assert kept["status"] == "error" and destination.read_bytes() == original_bytes
                listed = await success("list_paper_writing_projects", {"limit": 5, "offset": 0})
                assert listed["data"]["total"] == 1
                current = await success("get_paper_writing_project", {"project_id": project_id})
                assert current["data"]["revision"] == 4

                # Deliberately corrupt only this tmp_path fixture cache: retained
                # reading cards must not make stale SHA evidence usable again.
                from paperflow.engine.literature.store import PaperStore
                cache = PaperStore(str(paper_home)).pdf_path(row["file_sha256"])
                cache.write_bytes(cache.read_bytes() + b"\n% changed offline fixture\n")
                stale = await success("prepare_paper_manuscript", {"project_id": project_id})
                assert stale["data"]["evidence_total"] == 0 and stale["data"]["evidence_matrix"] == []
                invalid_output = tmp_path / "must-not-export-stale.docx"
                stale_export = await call("export_paper_manuscript", {"project_id": project_id, "output_path": str(invalid_output)})
                assert stale_export["status"] == "error" and not invalid_output.exists()
                stale_save = await call("save_paper_manuscript", {"project_id": project_id, "draft": draft,
                                                                  "expected_revision": 4, "change_note": "旧SHA引用必须拒绝"})
                assert stale_save["status"] == "error"
                assert not research_home.exists() and not journal_home.exists()

    asyncio.run(asyncio.wait_for(workflow(), timeout=90))
